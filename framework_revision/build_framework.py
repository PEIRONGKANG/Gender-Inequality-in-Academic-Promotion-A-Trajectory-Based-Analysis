#!/usr/bin/env python3
"""Build clean and tracked dissertation-framework DOCX files."""
from __future__ import annotations
import argparse, importlib.util, zipfile
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

BASE = Path(__file__).resolve().parents[1] / "proposal_revision" / "build_proposal.py"
spec = importlib.util.spec_from_file_location("proposal_builder", BASE)
b = importlib.util.module_from_spec(spec); spec.loader.exec_module(b)

RQ1 = "Where do gender differences appear in the observed academic promotion timeline in Japan?"
RQ2 = "How is the number of institutional moves associated with the promotion timeline, and does this association differ between women and men?"
RQ3 = "How is access to KAKEN-related academic resources associated with promotion across different career stages, and do these associations differ between women and men?"
OVERALL_RQ = "How do gender differences in Japanese academic careers emerge and accumulate through the promotion timeline, institutional mobility, and access to academic resources?"

def h(d, text, level=1): b.add_heading(d, text, level)
def p(d, text): b.add_para(d, text)
def placeholder(d, descriptive=True):
    h(d, "Results placeholder", 3)
    items=[]
    if descriptive: items.append(("Descriptive finding", "[Insert result here]"))
    items += [("Main model", "[Insert coefficient, uncertainty interval, and interpretation here]"),
              ("Gender comparison", "[Insert predicted probabilities or marginal effects here]"),
              ("Robustness", "[Insert robustness summary here]"),
              ("Substantive interpretation", "[Insert interpretation linked to theory here]")]
    for label, value in items:
        q=d.add_paragraph(); q.paragraph_format.space_after=Pt(2)
        r=q.add_run(label+": "); r.bold=True; q.add_run(value)

def add_page_field(paragraph):
    paragraph.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    run=paragraph.add_run(); begin=b.OxmlElement("w:fldChar"); begin.set(b.qn("w:fldCharType"),"begin")
    instr=b.OxmlElement("w:instrText"); instr.set(b.qn("xml:space"),"preserve"); instr.text=" PAGE "
    separate=b.OxmlElement("w:fldChar"); separate.set(b.qn("w:fldCharType"),"separate")
    text=b.OxmlElement("w:t"); text.text="1"; run._r.extend([begin,instr,separate,text,b.OxmlElement("w:fldChar")]); run._r[-1].set(b.qn("w:fldCharType"),"end")

def apply_tohoku_book_format(d):
    d.settings.odd_and_even_pages_header_footer=True
    for section in d.sections:
        section.page_width=Inches(7.25); section.page_height=Inches(10.5)
        section.left_margin=Inches(1.25); section.right_margin=Inches(1.25)
        section.top_margin=Inches(1.1); section.bottom_margin=Inches(1.1)
        section.header_distance=Inches(.45); section.footer_distance=Inches(.45)
        section.different_first_page_header_footer=True
        for header in (section.header,section.even_page_header,section.first_page_header):
            for p_ in list(header.paragraphs): p_._element.getparent().remove(p_._element)
        odd=section.header.add_paragraph(); add_page_field(odd)
        even=section.even_page_header.add_paragraph(); add_page_field(even); even.alignment=WD_ALIGN_PARAGRAPH.LEFT

def title(d, logo):
    q=d.add_paragraph(); q.alignment=WD_ALIGN_PARAGRAPH.CENTER; q.add_run().add_picture(str(logo),width=Inches(.9))
    for text,size,bold,color in [
        ("Graduate School of Arts and Letters",11,False,b.GRAY),("TOHOKU UNIVERSITY",16,True,b.BLUE),
        ("Identifying Gendered Patterns in Japanese Academic Careers",19,True,RGBColor(0,0,0)),
        ("Promotion Timeline, Institutional Mobility, and KAKEN-Related Academic Resources",14,False,b.GRAY),
        ("Doctoral Dissertation Framework",14,False,b.GRAY),("PEI Rongkang",14,True,RGBColor(0,0,0))]:
        q=d.add_paragraph(); q.alignment=WD_ALIGN_PARAGRAPH.CENTER; q.paragraph_format.space_after=Pt(6)
        r=q.add_run(text); r.font.name="Arial"; r.font.size=Pt(size); r.bold=bold; r.font.color.rgb=color
    p(d,"Dissertation submitted to"); d.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
    p(d,"Graduate School of Arts and Letters, Tohoku University"); d.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
    p(d,"in fulfillment of the requirements for the degree of"); d.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
    p(d,"Doctor of Philosophy in Art and Letters"); d.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
    table=d.add_table(rows=4,cols=2); table.cell(0,0).text="Supervisor:"; table.cell(0,1).text="Prof. LYU Zeyu"; table.cell(1,0).text="Referees:"; table.cell(1,1).text="Prof. __________________"; table.cell(2,1).text="Prof. __________________"; table.cell(3,1).text="Prof. __________________"; b.format_table(table,widths=[1400,3600],font_size=9)
    p(d,"Submission Date: July 2026"); d.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
    d.add_page_break()

def front(d):
    h(d,"Abstract")
    p(d,"Gender inequality in academic careers is visible not only in who reaches senior rank, but also in when advancement occurs and in the organizational and scholarly resources available along the way. This dissertation uses longitudinal researchmap career histories linked to KAKEN project records to examine academic careers in Japan. It brings the promotion timeline, institutional mobility, and access to academic resources into a common career-course framework while keeping the three empirical studies analytically distinct.")
    p(d,"Study 1 identifies where gender differences appear in the observed promotion timeline from the PhD to Associate Professor and from Associate Professor to Full Professor. Study 2 asks whether the number of institutional moves is associated with the promotion timeline and whether that relationship differs by gender. Sequence analysis provides a secondary account of recurring patterns of institutional stability and movement. Study 3 examines KAKEN-related academic resources across two career stages. It considers participation, project leadership, funding, and project-based network position before promotion.")
    p(d,"The dissertation contributes to research on gendered organizations and academic careers by treating the ability to move and the accumulation of research resources as career conditions that may be unequally distributed. Researcher-year event-history models preserve right-censoring and temporal order, while the sequence analysis describes longer career pathways. The evidence is descriptive and associational; the observational design does not establish causal effects.")
    h(d,"Acknowledgements"); p(d,"I am grateful to my supervisor, committee members, colleagues, and collaborators for their guidance. This section will be completed in the final dissertation.")
    h(d,"Table of Contents")
    for x in ["Chapter 1. Introduction","Chapter 2. Gender Inequality in Japanese Academic Careers","Chapter 3. Data and Methods","Chapter 4. Study 1: Gendered Promotion Timeline","Chapter 5. Study 2: Institutional Mobility and the Promotion Timeline","Chapter 6. Study 3: KAKEN Resources and Stage-Specific Promotion","Chapter 7. General Discussion","Chapter 8. Conclusion","Appendices","References"]:
        q=d.add_paragraph(x); q.paragraph_format.space_after=Pt(1)
    h(d,"List of Tables"); p(d,"Table 3.1. Core measures and analytical roles\nTable A.1. Variable construction and validation\n[Additional tables will be inserted as empirical chapters are completed]")
    h(d,"List of Figures"); p(d,"Figure 1.1. Dissertation logic\nFigure 4.1. Study 1 analytic samples\nFigure 4.2. Study 1 observed promotion timeline\n[Additional figures will be inserted as empirical chapters are completed]")
    h(d,"List of Abbreviations"); p(d,"AP — Associate Professor\nFP — Full Professor\nKAKEN — Grants-in-Aid for Scientific Research\nPI — Principal Investigator\nCo-I — Co-Investigator\nJST — Japan Science and Technology Agency")
    d.add_page_break()

def chapter1(d):
    d.add_page_break()
    h(d,"Chapter 1. Introduction")
    h(d,"1.1 Background",2); p(d,"Academic careers unfold through a series of appointments, promotions, institutional moves, and opportunities to secure research support. In Japan, women remain underrepresented at senior academic ranks even as participation in doctoral education and research has expanded. Cross-sectional rank distributions show the outcome of this process, but they reveal little about when inequalities emerge or how careers develop before promotion.")
    h(d,"1.2 Research Problem",2); p(d,"Three gaps guide the dissertation. The promotion timeline is often studied only among people whose transitions are observed. Institutional mobility is commonly treated as individual choice even though the capacity to move depends on organizational opportunity and private circumstances. Research funding is frequently measured as an outcome rather than as a career resource that may have different relevance at different stages.")
    h(d,"1.3 Research Questions",2)
    p(d,"Overall research question. "+OVERALL_RQ)
    for n,rq in [("RQ1",RQ1),("RQ2",RQ2),("RQ3",RQ3)]: p(d,f"{n}. {rq}")
    h(d,"1.4 Theoretical Orientation",2); p(d,"The argument draws on gendered organizations, linked lives, status attainment, and cumulative advantage. Organizational rules can appear neutral while rewarding career patterns that are easier for some researchers to sustain. Institutional mobility may open opportunities, but the ability to move may itself be an unequally distributed career resource. Likewise, access to academic resources and the career value attached to those resources may vary by gender and career stage.")
    h(d,"1.5 Data and Methodological Overview",2); p(d,"The dissertation links researchmap career records to KAKEN project data and organizes them by fiscal year. Study 1 compares the observed promotion timeline across career stages. Studies 2 and 3 use researcher-year risk sets that retain researchers who have not yet experienced promotion. Study 2 also uses multichannel sequence analysis to describe recurring mobility pathways.")
    h(d,"1.6 Dissertation Structure",2); p(d,"The empirical sequence is straightforward: promotion timeline, mobility, and academic resources. Chapter 4 establishes where gender differences appear. Chapter 5 examines institutional movement. Chapter 6 turns to KAKEN-related resources across career stages.")
    h(d,"1.7 Contribution",2); p(d,"The dissertation connects the promotion timeline to the organizational and scholarly conditions under which promotion occurs. It also extends completed-transition comparisons by retaining right-censored careers in the later studies.")

def chapter2(d):
    d.add_page_break()
    h(d,"Chapter 2. Gender Inequality in Japanese Academic Careers")
    sections=[
    ("2.1 Promotion Timeline: Locating Gender Differences","The first argument concerns where inequality appears in the academic career. Japanese careers combine rank progression with substantial variation across fields and institutions. Gender gaps in senior rank can arise because women and men differ in entry into promotion, in the promotion timeline among those promoted, or in whether a transition is observed at all. Study 1 separates these processes and identifies the career stage at which the observed difference is concentrated (Allison and Long, 1987; Pei et al., 2026)."),
    ("2.2 Institutional Mobility: The Organizational Pathway","Locating a difference does not explain the organizational path through which careers advance. Institutional moves can open access to positions and organizational settings, yet moving is not a costless individual choice. Gendered-organization theory directs attention to the rules and expectations embedded in academic employment, while linked-lives research shows that career decisions are coordinated with family and partners (Acker, 1990; Bielby and Bielby, 1992). The ability to move may itself be an unequally distributed career resource. Study 2 therefore tests whether cumulative move count is related to the promotion timeline and whether that relationship differs by gender."),
    ("2.3 Academic Resources: The Accumulation Pathway","Organizational movement is only one route through which careers develop. KAKEN participation, project leadership, funding, and collaboration position represent resources accumulated within and across institutions. Cumulative-advantage theory suggests that early access can shape later opportunities, but the value attached to the same resource may vary by career stage and gender (Merton, 1968; DiPrete and Eirich, 2006; Bol et al., 2018). Study 3 compares these associations in the early-career-to-AP and AP-to-FP transitions."),
    ("2.4 An Integrated Gendered Career Process","The three arguments form a sequence rather than three parallel topics. Study 1 locates the observed difference. Study 2 examines organizational movement as one pathway connected to that difference. Study 3 examines the accumulation of academic resources across career stages. Gendered organizations provide the common explanation for why mobility opportunities and the career value of resources may not be distributed or evaluated uniformly."),
    ("2.5 Research Gap and Analytical Expectations","Existing work rarely examines the promotion timeline, institutional movement, and project-based academic resources with the same longitudinal population. The dissertation links them while preserving distinct outcomes and risk sets. It asks where gender differences appear, whether organizational movement is associated with the observed timeline, and whether KAKEN-related resources are differently associated with promotion across career stages.")]
    for title_,text in sections: h(d,title_,2); p(d,text)

def chapter3(d):
    d.add_page_break()
    h(d,"Chapter 3. Data and Methods")
    sections=[
    ("3.1 Data Sources","researchmap provides career histories, education, affiliations, and researcher profiles (JST, 2026). KAKEN records provide project membership, roles, project categories, and annual budgets (JSPS, 2026a, 2026b). Fiscal year is the common time unit."),
    ("3.2 Research Population","The population requires a valid researcher identifier, a plausible PhD year, usable career history, gender for the primary comparison, and field information. Researchers without observed promotion remain in the later risk sets when observation coverage is adequate."),
    ("3.3 Researcher and Institution Linkage","Linkage distinguishes the KAKEN website identifier from the JPR researchmap user identifier used for the dissertation cohort. Unlinked project members remain in team-size denominators. Institution aliases are harmonized before institutional moves are counted."),
    ("3.4 Career and Promotion Measures","Promotion is the first validated appointment to Associate Professor or Full Professor. Japanese and English title patterns are audited. Time since PhD defines the main analysis clock, and post-event years are excluded from each risk set."),
    ("3.5 Institutional Mobility Measures","An institutional move is a change in canonical primary institution between consecutive observed records. The primary measure is cumulative move count before year t. Direction, internal versus external promotion, time since last move, prestige, and institution size are secondary measures."),
    ("3.6 KAKEN Resource Measures","KAKEN resources are organized into access, project leadership, funding, and project-based network position. Funding is measured using full exposure and fractional allocation. The original project category is retained."),
    ("3.7 Event-History Analysis","Discrete-time event-history models estimate promotion risk with a flexible baseline hazard. Predictors are measured before year t. Complementary log-log models are primary; logistic and Cox specifications provide comparisons."),
    ("3.8 Sequence Analysis","Sequence analysis will identify recurring patterns of institutional stability and movement (Abbott and Tsay, 2000; Gauthier et al., 2010). Channel construction, distance alternatives, clustering diagnostics, and sensitivity windows are documented in the Appendix."),
    ("3.9 Missing Data and Validation","No KAKEN record is distinguished from failed linkage, and missing career history is not coded as career exit. Sample-flow, missingness, rank-title, and linkage audits accompany each empirical chapter."),
    ("3.10 Ethics and Limitations","The analysis uses structured records of professional careers and reports only aggregate results. Temporal ordering improves interpretation but does not remove selection, measurement error, or unobserved confounding.")]
    for title_,text in sections: h(d,title_,2); p(d,text)

def chapter4(d,media):
    d.add_page_break()
    h(d,"Chapter 4. Study 1: Gendered Promotion Timeline")
    h(d,"4.1 Introduction",2); p(d,"Study 1 provides the diagnostic starting point by locating gender differences across two observed promotion transitions.")
    h(d,"4.2 Research Question",2); p(d,"RQ1. "+RQ1)
    h(d,"4.3 Data and Measures",2); p(d,"The completed study uses 74,344 discipline-mapped PhD holders. The observed-transition samples contain 20,808 researchers for PhD to AP and 10,747 for AP to FP; positive-duration samples contain 18,794 and 10,416 researchers, respectively.")
    b.add_picture(d,media/"image3.png",6.5); b.add_caption(d,"Figure 4.1. Study 1 official analytic and observed-transition samples. Source: Pei et al. (2026).")
    h(d,"4.4 Analytical Strategy",2); p(d,"The promotion timeline is measured between validated milestones. Completed-transition comparisons are interpreted alongside survival-based checks that retain right-censoring.")
    h(d,"4.5 Results",2); p(d,"The completed analysis identifies the PhD-to-Associate-Professor transition as the main site of observed heterogeneity. Among positive-duration completed transitions, mean timelines are 7.83 years for women and 8.77 years for men from PhD to AP, and 8.22 years for women and 8.60 years for men from AP to FP. These conditional means do not imply that women have a higher probability of promotion.")
    b.add_picture(d,media/"image4.png",6.5); b.add_caption(d,"Figure 4.2. Study 1 mean observed promotion timeline by gender and career stage. Source: Pei et al. (2026).")
    h(d,"4.6 Discussion",2); p(d,"The early-career transition deserves closer attention, but the promotion timeline alone cannot reveal how organizational opportunities shape advancement. The distinction between transition duration and entry into promotion motivates the risk-set designs that follow.")
    h(d,"4.7 Chapter Summary",2); p(d,"Study 1 shows where differences in the observed promotion timeline appear. Chapter 5 asks whether institutional movement is related to that timeline.")

def chapter5(d):
    d.add_page_break()
    h(d,"Chapter 5. Study 2: Institutional Mobility and the Promotion Timeline")
    h(d,"5.1 Introduction",2); p(d,"Institutional mobility may offer career opportunities, but the capacity to move is not equally distributed. Study 2 therefore centers on the number of institutional moves rather than treating every feature of mobility as a separate question.")
    h(d,"5.2 Theoretical Expectations",2); p(d,"H2.1 More institutional moves are associated with the promotion timeline.\nH2.2 The association between move count and the promotion timeline differs between women and men.\nH2.3 Upward moves are more strongly associated with promotion than lateral or downward moves.")
    h(d,"5.3 Research Question",2); p(d,"RQ2. "+RQ2)
    h(d,"5.4 Variables and Measures",2); p(d,"The variable hierarchy is fixed in advance. The primary explanatory variable is cumulative institutional move count before year t. It provides the clearest test of whether repeated organizational movement is related to the promotion timeline. Secondary mobility variables distinguish upward, lateral, and downward moves, internal and external promotion, time since the last move, and mobility per observed career year. Institutional ranking and size are contextual measures used to describe destinations and test robustness; they are not competing explanations or substitutes for move count.")
    h(d,"5.5 Analytical Strategy",2); p(d,"The core event-history analysis models time to first Associate Professor appointment using move count and a gender-by-move-count interaction. Directional measures enter secondary models after the main association has been established. Ranking and size measures are introduced only in contextual or robustness specifications. Multichannel sequence analysis is supplementary: it identifies recurring patterns of institutional stability and movement but does not define the main outcome or replace the event-history model. Where coverage permits, the AP-to-FP transition is examined separately.")
    h(d,"5.6 Results",2); placeholder(d)
    h(d,"5.7 Robustness Checks",2); p(d,"[Insert comparison of move definitions, prestige measures, observation windows, and high-coverage samples here]")
    h(d,"5.8 Discussion",2); p(d,"[Insert discussion of mobility as opportunity and constraint, linked to gendered organizations and linked lives]")
    h(d,"5.9 Chapter Summary",2); p(d,"[Insert chapter summary here] Mobility captures movement between organizations; Chapter 6 examines access to academic resources within and across career stages.")

def chapter6(d):
    d.add_page_break()
    h(d,"Chapter 6. Study 3: KAKEN Resources and Stage-Specific Promotion")
    h(d,"6.1 Introduction",2); p(d,"Study 3 examines whether access to KAKEN-related academic resources is associated with promotion and whether the association changes across career stages. It contains two core promotion analyses: the early-career transition to Associate Professor and the later transition from Associate Professor to Full Professor. A descriptive analysis around AP appointment is supplementary and does not replace the second promotion model.")
    h(d,"6.2 Theoretical Expectations",2); p(d,"H3.1 Greater access to KAKEN-related academic resources is associated with faster promotion.\nH3.2 The association between KAKEN resources and promotion differs between women and men.\nH3.3 The association between KAKEN resources and promotion varies across career stages.")
    h(d,"6.3 Research Question",2); p(d,"RQ3. "+RQ3)
    h(d,"6.4 KAKEN Resource Measures",2); p(d,"Resources are grouped into participation, project leadership, funding, and project-based network position. Primary measures include any KAKEN participation, PI and Co-I roles, direct funding, degree, PI outdegree, and PI indegree. Weighted degree is reserved for robustness analysis. All predictors for promotion in fiscal year t are constructed from information observed before t.")
    h(d,"6.5 Stage 1: PhD/Early Career to Associate Professor",2); p(d,"The primary Stage 1 risk set begins in the first fiscal year after PhD completion. Researchers remain at risk until their first validated Associate Professor appointment or the observation endpoint. Researchers who are not observed as AP remain as right-censored cases, and post-event person-years are excluded.")
    p(d,"Where postdoctoral appointments can be identified reliably, postdoctoral status is constructed as a time-varying career state. A secondary Postdoc-to-Associate-Professor analysis will examine promotion after entry into an observed postdoctoral appointment. An observed postdoc is not an inclusion requirement for the primary Stage 1 analysis: some researchers move directly into faculty or other academic positions, and postdoctoral histories may be incomplete.")
    placeholder(d)
    h(d,"6.6 Stage 2: Associate Professor to Full Professor",2); p(d,"The Stage 2 risk set begins in the first fiscal year after the validated Associate Professor appointment and ends at first observed Full Professor appointment or censoring. Researchers already observed as Full Professor when the AP appointment year first becomes observable are excluded or flagged for chronology review.")
    p(d,"KAKEN predictors in this stage are measured after the AP appointment and before the FP transition year. Pre-AP resources may be retained as baseline career history, but they are not substituted for the post-AP resource measures that define the core Stage 2 analysis. This separation ensures that the stage comparison concerns resources accumulated while researchers are at risk for the relevant promotion.")
    placeholder(d)
    h(d,"6.7 Gender Differences Across Career Stages",2); p(d,"Separate stage-specific model families are primary. Gender-by-resource terms are estimated within each stage, and stage-by-resource comparisons are made across aligned specifications. A pooled model with gender-by-stage-by-resource interactions will be used only when the data provide adequate events and common support."); placeholder(d,False)
    h(d,"6.8 Supplementary Post-AP Resource Trajectories",2); p(d,"A descriptive event-time analysis may trace KAKEN participation, PI status, direct funding, and network position before and after Associate Professor appointment. Its purpose is to show how observed resources change around AP appointment. It is not a causal design and does not replace the Associate-Professor-to-Full-Professor promotion model.")
    p(d,"[Insert descriptive post-AP event-time trajectories here]")
    h(d,"6.9 Robustness Checks",2); p(d,"[Insert comparison of lag windows, funding allocations, linkage-quality samples, and alternative network measures here]")
    h(d,"6.10 Discussion",2); p(d,"[Insert discussion of cumulative advantage, project leadership, and gendered returns to resources here]")
    h(d,"6.11 Chapter Summary",2); p(d,"[Insert chapter summary and transition to the General Discussion here]")

def ending(d):
    d.add_page_break()
    h(d,"Chapter 7. General Discussion")
    for s in ["7.1 Summary of Findings","7.2 Promotion Timeline","7.3 Mobility as a Gendered Career Strategy","7.4 Gendered Access to Academic Resources","7.5 Career-Stage Differences","7.6 Theoretical Contributions","7.7 Policy Implications","7.8 Limitations","7.9 Future Research"]:
        h(d,s,2); p(d,"[Insert synthesis after the relevant empirical chapters are completed]")
    d.add_page_break(); h(d,"Chapter 8. Conclusion"); p(d,"[Insert concise conclusion after all empirical chapters and the General Discussion are complete]")
    h(d,"Appendices")
    for s,t in [
      ("Appendix A. Variable Dictionary","[Insert final variable dictionary and provenance table]"),
      ("Appendix B. Institutional Ranking and Size Construction","THE, QS, and ARWU measures remain separate. Annual percentiles, pre-move medians, coverage rules, and institution-size measures will be documented here."),
      ("Appendix C. Mobility Coding","[Insert primary-affiliation hierarchy, move-event rules, direction thresholds, and internal/external promotion validation]"),
      ("Appendix D. KAKEN and Network Construction","The raw KAKEN project-category field, linkage-status definitions, budget-allocation denominators, role harmonization, and project-year edge construction are documented here."),
      ("Appendix E. Robustness Specifications","[Insert alternative lag windows, baseline hazards, sequence distances, clustering diagnostics, and sensitivity samples]"),
      ("Appendix F. Data-Quality Diagnostics","[Insert sample flow, missingness, title audit, linkage audit, and coverage diagnostics]"),
      ("Appendix G. Additional Tables and Figures","[Insert supplementary tables and figures]")]: h(d,s,2); p(d,t)
    b.add_references(d)

def build(source,out,work,school_template):
    d=Document(source); b.clear_body(d); b.configure_styles(d)
    media=work/"media"; media.mkdir(parents=True,exist_ok=True)
    with zipfile.ZipFile(source) as z:
        for n in z.namelist():
            if n.startswith("word/media/") and not n.endswith("/"): (media/Path(n).name).write_bytes(z.read(n))
    with zipfile.ZipFile(school_template) as z:
        school_logo=media/"tohoku_template_logo.png"; school_logo.write_bytes(z.read("logo.png"))
    title(d,school_logo); front(d); chapter1(d); chapter2(d); chapter3(d); chapter4(d,media); chapter5(d); chapter6(d); ending(d); apply_tohoku_book_format(d)
    out.parent.mkdir(parents=True,exist_ok=True)
    raw=out.with_suffix(".raw.docx"); d.save(raw); b.strip_comments(raw,out); raw.unlink()

def main():
    a=argparse.ArgumentParser(); a.add_argument("--source",type=Path,required=True); a.add_argument("--school-template",type=Path,required=True); a.add_argument("--out-dir",type=Path,required=True); a.add_argument("--work-dir",type=Path,required=True); x=a.parse_args()
    source=x.source.resolve(); out=x.out_dir.resolve(); work=x.work_dir.resolve(); work.mkdir(parents=True,exist_ok=True)
    clean=out/"Doctoral_Dissertation_Framework_PEI_Rongkang_clean.docx"; tracked=out/"Doctoral_Dissertation_Framework_PEI_Rongkang_tracked.docx"
    build(source,clean,work,x.school_template.resolve()); b.make_tracked(clean,tracked); print(clean); print(tracked)
if __name__=="__main__": main()
