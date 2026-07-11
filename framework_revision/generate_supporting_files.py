#!/usr/bin/env python3
import csv
from pathlib import Path

ROOT=Path(__file__).resolve().parent
def write(name,header,rows):
    with (ROOT/name).open("w",newline="",encoding="utf-8-sig") as f:
        w=csv.writer(f); w.writerow(header); w.writerows(rows)

write("research_question_mapping.csv",["study","formal_research_question","chapter","former_questions_relocated_to"],[
 ["Study 1","Where do gender differences appear in observed academic promotion timelines in Japan?","Chapter 4","Stage and discipline comparisons in measures and analysis"],
 ["Study 2","How is the number of institutional moves associated with promotion timing, and does this association differ between women and men?","Chapter 5","H2.1-H2.3; secondary mobility measures; sequence analysis; robustness checks"],
 ["Study 3","How is access to KAKEN-related academic resources associated with promotion across different career stages, and do these associations differ between women and men?","Chapter 6","H3.1-H3.3; stage-specific models; selected interactions; robustness checks"],
])
terms=[
 ("researchmap","lowercase except at sentence start","Researcher database"),("KAKEN","uppercase","Funding system"),
 ("Associate Professor","title case; AP after definition","Academic rank"),("Full Professor","title case; FP after definition","Academic rank"),
 ("PI","PI","Project-specific Principal Investigator role"),("Co-I","hyphenated","Project-specific Co-Investigator role"),
 ("institutional mobility","preferred broad term","Changes between institutions"),("institutional move","countable event","Change in canonical primary institution"),
 ("promotion timing","preferred general term","Elapsed timing of promotion"),("promotion risk","event-history interpretation only","Conditional probability/hazard"),
 ("researcher-year","hyphenated","Panel unit"),("event-history analysis","hyphenated","Risk-set model family"),
 ("multichannel sequence analysis","full term","Secondary pathway analysis"),("KAKEN-related academic resources","hyphenated","Access, leadership, funding, and network position"),
 ("project-based collaboration networks","hyphenated","Networks derived from shared KAKEN projects"),("career stage","two words","PhD-to-AP or AP-to-FP"),
 ("right-censoring","hyphenated","Event not observed by endpoint"),("time-ordered association","hyphenated","Predictor precedes outcome year")]
write("terminology_control_table.csv",["preferred_term","usage_rule","definition"],terms)

source=ROOT.parent/"proposal_revision"/"citation_audit.csv"
rows=[]
from docx import Document
doc=Document(ROOT/"Doctoral_Dissertation_Framework_PEI_Rongkang_clean.docx")
text="\n".join(p.text for p in doc.paragraphs)
chapter_text=text.rsplit("\nReferences\n",1)[0]
with source.open(encoding="utf-8-sig",newline="") as f:
    for r in csv.DictReader(f):
        citation=f"{r.get('authors','')} ({r.get('year','')})"
        surname=r.get("authors","").split(",",1)[0].strip()
        used="yes" if surname and surname in chapter_text and r.get("year","") in chapter_text else "no"
        note=r.get("notes","") or r.get("exact_claim_supported","")
        if used=="no": note=(note+"; retained as verified framework bibliography but not yet cited in chapter prose").strip("; ")
        rows.append([citation,r.get("verified","yes"),r.get("doi",""),used,"yes",note])
write("citation_audit.csv",["citation","verified","doi","used_in_text","in_reference_list","notes"],rows)

(ROOT/"revision_memo.md").write_text("""# Revision Memo

## Dissertation conversion

The proposal has been restructured as an eight-chapter dissertation framework. Proposal milestones and workflow catalogues were removed from the main body. Front matter now includes a dissertation-style contents list, lists of tables and figures, and abbreviations. The source DOCX was not overwritten.

## Research questions

The former RQ2a-RQ2e and RQ3a-RQ3d labels were removed. Each study now has one formal question. Their useful analytical content was retained as hypotheses, secondary measures, model choices, or robustness checks.

## Study 2

Study 2 now centers on cumulative institutional move count before year t and promotion timing. Move direction, internal versus external promotion, prestige, size, and sequence analysis are secondary. The chapter emphasizes that mobility can be both opportunity and constraint, and that the capacity to move is unequally distributed.

## Study 3

Study 3 groups KAKEN-related academic resources into access, project leadership, funding, and project-based network position. Separate model families cover PhD-to-AP and AP-to-FP promotion. Interactions are selective rather than exhaustive.

## Humanisation

Repeated causal disclaimers, symmetrical variable catalogues, generic transitions, status-code lists, and engineering-style workflow prose were reduced. The main text uses shorter arguments and stronger theoretical judgment. Raw field names, detailed linkage states, ranking formulas, sequence-distance alternatives, and diagnostic inventories were moved to the Appendices.

## Unfinished empirical sections

Study 1 retains only the previously reported sample counts and promotion-timeline results. Study 2, Study 3, the General Discussion, and the Conclusion contain explicit structured placeholders. No empirical values were created for unfinished analyses.
""",encoding="utf-8")

(ROOT/"unresolved_issues.md").write_text("""# Unresolved Issues

1. Study 2 descriptive, event-history, interaction, sequence, robustness, and discussion sections await empirical results.
2. Study 3 requires separate AP and FP risk-set analyses; no Stage 2 results are currently available.
3. General Discussion and Conclusion must be completed only after the empirical chapters are finalized.
4. Institution prestige and size coverage require validation before substantive interpretation.
5. Official institution properties should replace temporary affiliation-string heuristics when the institution property table becomes available.
6. KAKEN non-participation must remain distinct from unresolved linkage and limited source coverage.
7. Family and dual-career constraints are theoretically relevant but not directly observed.
8. The final dissertation should replace static front-matter lists with Word-generated page-numbered lists after chapter pagination stabilizes.
9. Citation records are inherited from the verified proposal audit; any new literature added later requires a fresh Zotero or authoritative-source check.
""",encoding="utf-8")
