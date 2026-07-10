#!/usr/bin/env python3
"""Build clean and tracked dissertation-proposal DOCX deliverables.

The source DOCX is used only as a style and institutional-identity template.
The script never writes to the source path.
"""

from __future__ import annotations

import argparse
import copy
import datetime as dt
import json
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from lxml import etree
from PIL import Image, ImageDraw, ImageFont

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W_NS}
BLUE = RGBColor(0x00, 0x72, 0xB2)
GRAY = RGBColor(0x4D, 0x57, 0x61)
CJK_IMAGE_PATH: Path | None = None


def clear_body(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_twips: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_twips))
    tc_w.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def format_table(table, widths=None, font_size=8.0) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    if widths:
        for row in table.rows:
            for idx, cell in enumerate(row.cells):
                set_cell_width(cell, widths[idx])
    set_repeat_table_header(table.rows[0])
    for r_idx, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if r_idx == 0:
                set_cell_shading(cell, "D9EAF7")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(1)
                paragraph.paragraph_format.space_after = Pt(1)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(font_size)
                    if r_idx == 0:
                        run.bold = True


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10.5)
    normal._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Hiragino Mincho ProN")
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(5)
    for name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11)):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLUE
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Hiragino Kaku Gothic ProN")
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(4)
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)


def add_para(doc: Document, text: str, *, bold_lead: str | None = None) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.widow_control = True
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    elif "准教授" in text:
        before, after = text.split("准教授", 1)
        p.add_run(before)
        cjk = p.add_run()
        if CJK_IMAGE_PATH is None:
            raise RuntimeError("CJK title-pattern image has not been initialized")
        cjk.add_picture(str(CJK_IMAGE_PATH), width=Inches(0.44))
        doc_pr = cjk._element.find(".//wp:docPr", namespaces={"wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"})
        if doc_pr is not None:
            doc_pr.set("descr", "准教授")
        p.add_run(after)
    else:
        p.add_run(text)
    for run in p.runs:
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Hiragino Mincho ProN")


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        p.add_run(item)


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = False
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(9)


def add_picture(doc: Document, path: Path, width: float) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(width))


def _fonts():
    regular = "/System/Library/Fonts/Supplemental/Arial.ttf"
    bold = "/System/Library/Fonts/Supplemental/Arial Bold.ttf"
    return (ImageFont.truetype(regular, 34), ImageFont.truetype(bold, 42),
            ImageFont.truetype(regular, 27), ImageFont.truetype(bold, 31))


def _center_multiline(draw, xy, text, font, fill="#222222", spacing=8):
    draw.multiline_text(xy, text, font=font, fill=fill, anchor="mm", align="center", spacing=spacing)


def _arrow(draw, start, end, fill="#5D6872", width=7):
    draw.line([start, end], fill=fill, width=width)
    x, y = end
    draw.polygon([(x, y), (x - 22, y - 14), (x - 22, y + 14)], fill=fill)


def make_framework_figure(path: Path) -> None:
    img = Image.new("RGB", (2160, 650), "white")
    draw = ImageDraw.Draw(img)
    regular, bold, small, small_bold = _fonts()
    boxes = [(80, 170, 620, 480), (810, 170, 1350, 480), (1540, 170, 2080, 480)]
    titles = ["Study 1: Timing", "Study 2: Mobility pathways", "Study 3: Resources and returns"]
    bodies = ["Where do gender differences\nappear in promotion timelines?",
              "Who moves, in what direction,\nand through which institutions?",
              "Who accesses KAKEN resources\nand project-based networks?"]
    fills = ["#DCEAF7", "#DDEFE9", "#FFF2C6"]
    for i, box in enumerate(boxes):
        draw.rounded_rectangle(box, radius=22, fill=fills[i], outline="#0072B2", width=5)
        cx = (box[0] + box[2]) // 2
        _center_multiline(draw, (cx, 235), titles[i], small_bold, fill="#006DAA")
        _center_multiline(draw, (cx, 350), bodies[i], small)
        if i < 2:
            _arrow(draw, (box[2] + 28, 325), (boxes[i + 1][0] - 28, 325))
    _center_multiline(draw, (1080, 575),
                      "timing  >  mobility pathways  >  academic resources and returns", regular)
    img.save(path)


def make_mobility_figure(path: Path) -> None:
    img = Image.new("RGB", (2160, 620), "white")
    draw = ImageDraw.Draw(img)
    regular, bold, small, small_bold = _fonts()
    labels = ["Career history\nspells", "Canonical primary\ninstitution-year",
              "Institutional move\nevents", "Risk panel +\nsequence channels"]
    fills = ["#F2F2F2", "#DCEAF7", "#DDEFE9", "#FFF2C6"]
    x_positions = [70, 600, 1130, 1660]
    for i, x in enumerate(x_positions):
        box = (x, 150, x + 430, 410)
        draw.rounded_rectangle(box, radius=20, fill=fills[i], outline="#0072B2", width=5)
        _center_multiline(draw, (x + 215, 280), labels[i], small_bold)
        if i < 3:
            _arrow(draw, (x + 445, 280), (x_positions[i + 1] - 15, 280))
    _center_multiline(draw, (1080, 525),
                      "Move direction uses lagged prestige and size measures; geography enters only after validated crosswalks.", small)
    img.save(path)


def make_kaken_timing_figure(path: Path) -> None:
    img = Image.new("RGB", (2160, 620), "white")
    draw = ImageDraw.Draw(img)
    regular, bold, small, small_bold = _fonts()
    y = 340
    xs = [180, 500, 820, 1140, 1460, 1860]
    draw.line([(xs[0], y), (xs[-1], y)], fill="#5D6872", width=8)
    labels = ["t-5", "t-4", "t-3", "t-2", "t-1", "year t"]
    for i, x in enumerate(xs):
        color = "#E69F00" if i == 5 else "#0072B2"
        draw.ellipse((x - 25, y - 25, x + 25, y + 25), fill=color)
        _center_multiline(draw, (x, y + 75), labels[i], small_bold)
    draw.rounded_rectangle((170, 70, 1540, 235), radius=20, fill="#DCEAF7", outline="#0072B2", width=5)
    _center_multiline(draw, (855, 152), "Lagged KAKEN access, roles, funding, category, and network position", small_bold)
    draw.line([(1460, 235), (1460, y - 35)], fill="#5D6872", width=6)
    draw.rounded_rectangle((1650, 70, 2070, 235), radius=20, fill="#FFF2C6", outline="#0072B2", width=5)
    _center_multiline(draw, (1860, 152), "First observed AP\ntransition", small_bold)
    draw.line([(1860, 235), (1860, y - 35)], fill="#5D6872", width=6)
    _center_multiline(draw, (1080, 545),
                      "Primary: t-1 and past3 (t-3 through t-1); robustness: past4, past5, cumulative prior exposure", small)
    img.save(path)


def make_cjk_title_pattern(path: Path) -> None:
    img = Image.new("RGBA", (220, 68), (255, 255, 255, 0))
    draw = ImageDraw.Draw(img)
    font = ImageFont.truetype("/System/Library/Fonts/ヒラギノ明朝 ProN.ttc", 52, index=0)
    draw.text((4, 31), "准教授", font=font, fill="#111111", anchor="lm")
    img.save(path)


def add_title_page(doc: Document, logo: Path) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(logo), width=Inches(0.9))
    for text, size, bold, color in [
        ("Graduate School of Arts and Letters", 12, False, GRAY),
        ("TOHOKU UNIVERSITY", 16, True, BLUE),
        ("Identifying Gendered Patterns in Japanese Academic Careers: Promotion Timelines, Institutional Mobility, KAKEN Resources, and Collaboration Networks", 18, True, RGBColor(0, 0, 0)),
        ("Doctoral Dissertation Proposal", 14, False, GRAY),
        ("PEI Rongkang", 14, True, RGBColor(0, 0, 0)),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(size)
        r.bold = bold
        r.font.color.rgb = color
    add_para(doc, "Proposal submitted to")
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for text in [
        "Graduate School of Arts and Letters, Tohoku University",
        "in partial fulfillment of the requirements for the degree of",
        "Doctor of Philosophy in Art and Letters",
    ]:
        p = doc.add_paragraph(text)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table = doc.add_table(rows=4, cols=2)
    table.cell(0, 0).text = "Supervisor:"
    table.cell(0, 1).text = "Prof. LYU Zeyu"
    table.cell(1, 0).text = "Referees:"
    table.cell(1, 1).text = "Prof. __________________________"
    table.cell(2, 1).text = "Prof. __________________________"
    table.cell(3, 1).text = "Prof. __________________________"
    format_table(table, widths=[1400, 4200], font_size=9)
    p = doc.add_paragraph("Submission Date: July 11, 2026")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()


def add_front_matter(doc: Document) -> None:
    add_heading(doc, "Abstract", 1)
    add_para(doc, "This dissertation examines where gendered differences appear in Japanese academic career trajectories and how those differences are associated with institutional mobility, KAKEN resources, project-based collaboration networks, and post-promotion trajectories. It links longitudinal researchmap career histories with Grants-in-Aid for Scientific Research (KAKEN) project records in a computational social science design. The dissertation is explicitly descriptive and associational: temporal ordering strengthens interpretation, but the observational data do not identify causal effects.")
    add_para(doc, "Study 1 is the completed diagnostic study of gendered promotion timelines. It compares the PhD-to-Associate-Professor and Associate-Professor-to-Full-Professor transitions while distinguishing observed transition duration, entry into the transition, and right-censoring. The completed evidence identifies the early-career transition as the main site of heterogeneity and shows why shorter durations among observed promotees cannot be interpreted as a higher probability of promotion.")
    add_para(doc, "Study 2 reframes institutional mobility as a potentially valuable but costly career strategy. It constructs institutional moves from consecutive career spells, distinguishes internal from external promotion where observable, and measures move frequency, timing, direction, prestige change, organizational-size change, and institutional stability. Researcher-year event-history models examine time-ordered mobility-promotion associations, while multichannel sequence analysis identifies typical combinations of rank, prestige-tier, and mobility states. The theoretical claim is that the ability to move may itself be an unequally distributed career resource; family and relocation costs are hypothesized mechanisms rather than directly observed variables.")
    add_para(doc, "Study 3 examines gender differences in observed access to KAKEN participation, PI and Co-I roles, funding intensity, original KAKEN project category, and KAKEN-based collaboration networks. A researcher-year promotion-risk design uses only t−1 or earlier information to assess whether pre-promotion resources are associated with first observed Associate Professor transition in year t. A secondary event-time extension describes KAKEN and network trajectories before and after Associate Professor appointment without treating promotion as an exogenous intervention. Together, the studies follow the logic timing → mobility pathways → academic resources and returns.")
    add_heading(doc, "Acknowledgement", 1)
    add_para(doc, "I am grateful to my supervisor, committee members, colleagues, and collaborators for their guidance on academic careers, gender inequality, longitudinal data construction, and computational social science. Remaining errors are my own.")
    add_heading(doc, "Table of Contents", 1)
    for line in [
        "Abstract", "Acknowledgement", "1. Introduction", "1.1 Research Background",
        "1.2 Research Questions", "1.3 Dissertation Structure",
        "2. Literature Review and Theoretical Framework",
        "2.1 Gender Inequality in Japanese Academic Careers",
        "2.2 Meritocracy, Gendered Evaluation, and Career-Relevant Academic Resources",
        "2.3 Career Mobility, Status Attainment, and Segmented Academic Labour Markets",
        "2.4 Gendered Organizations, Linked Lives, and Spatial Constraints",
        "2.5 Institutional Prestige, Organizational Scale, and Academic Mobility",
        "2.6 KAKEN Resources and Project Roles as Career-Relevant Academic Resources",
        "2.7 Project-Based Collaboration Networks and Academic Careers",
        "2.8 Cumulative Advantage and Post-Promotion Trajectories",
        "2.9 Longitudinal, Event-History, and Sequence Approaches",
        "3. Study 1: Gendered Promotion Timelines in Japanese Academia",
        "4. Study 2: Institutional Mobility as a Gendered Career Strategy",
        "5. Study 3: Gendered Access to and Returns from KAKEN Resources and Project-Based Collaboration Networks",
        "6. Data and Methodological Strategy", "7. Expected Contributions and Work Plan",
        "Appendix A. Variable Construction and Validation", "References",
    ]:
        p = doc.add_paragraph(line)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 0.9
        for run in p.runs:
            run.font.size = Pt(8.5)
    doc.add_page_break()


def add_intro(doc: Document, figs: Path) -> None:
    add_heading(doc, "1. Introduction", 1)
    add_heading(doc, "1.1 Research Background", 2)
    add_para(doc, "Academic careers are organized through temporally ordered transitions among training, appointments, institutions, funding roles, and collaborative relationships. Formal evaluation systems present these transitions as meritocratic, yet organizational rules and apparently neutral standards can reward career patterns that are easier for some researchers to sustain than for others (Acker, 1990; Nielsen, 2016; Schimanski & Alperin, 2018). A longitudinal design is therefore needed to distinguish when gender differences appear, which organizational pathways precede advancement, and how academic resources are distributed across career stages.")
    add_para(doc, "Japan is a consequential setting for this inquiry. Women remain a minority of the national research workforce, accounting for 18.5% of researchers in 2024, and field and rank distributions remain uneven (Statistics Bureau of Japan, 2025). Research on Japanese economics, medicine, and international engagement documents gendered career differences but does not provide a unified account of promotion timing, institutional movement, and access to project-based resources (Takahashi & Takahashi, 2015; Nagano et al., 2022; Watanabe, 2025). researchmap, operated by the Japan Science and Technology Agency (JST), provides broad career-history coverage, while KAKEN records make formal project participation, roles, categories, and budgets observable (JST, 2026; JSPS, 2026a). Their linkage enables a fiscal-year-level account that is unavailable from cross-sectional staff statistics alone.")
    add_para(doc, "The central gap is threefold. Existing evidence describes promotion timing, but less is known about the organizational pathways through which researchers move. Mobility may create access to better matches or higher-status institutions, yet moving also requires the capacity to absorb uncertainty, coordination, and relocation costs (Allison & Long, 1987; Fernández-Zubieta et al., 2016). Less is also known about whether KAKEN access, project leadership, funding intensity, and project-based relational positions are distributed differently by gender and associated with subsequent promotion. Finally, linked longitudinal evidence that separates pre-promotion predictors from promotion-year and post-promotion information remains limited.")
    add_para(doc, "This dissertation addresses these gaps through three empirical studies. The contribution is not a claim that mobility or KAKEN resources cause promotion. Instead, the design establishes transparent temporal order, retains right-censored researchers, distinguishes missing linkage from observed non-participation, and evaluates gender differences in observed pathways and time-ordered associations.")
    add_heading(doc, "1.2 Research Questions", 2)
    add_para(doc, "The overarching research question is: Where do gendered differences appear in Japanese academic career trajectories, and how are these differences associated with institutional mobility, KAKEN resources, project-based collaboration networks, and post-promotion trajectories?")
    add_para(doc, "RQ1. Where do gender differences appear in observed promotion timelines from PhD to Associate Professor and from Associate Professor to Full Professor, and how do right-censoring and selection into observed transitions qualify those comparisons?")
    for rq in [
        "RQ2a. How do women and men differ in the frequency, timing, and type of institutional mobility across academic career stages?",
        "RQ2b. Are upward, lateral, and downward institutional moves differently associated with PhD-to-Associate-Professor promotion, promotion timing, and final observed career position?",
        "RQ2c. Do the associations between institutional mobility and career outcomes differ by gender?",
        "RQ2d. What typical institutional-mobility pathways characterize Japanese academic careers, and how are women and men distributed across these pathways?",
        "RQ2e. Among researchers who reach similar academic destinations, do women and men arrive through different combinations of institutional stability, institutional mobility, prestige change, and organizational-size change?",
        "RQ3a. Do women and men differ in observed access to KAKEN participation, PI and Co-I roles, funding intensity, original KAKEN project category, and project-based network positions?",
        "RQ3b. Among researchers at risk of first Associate Professor appointment, are KAKEN resources and KAKEN-based network positions measured before year t associated with promotion in year t?",
        "RQ3c. Do the estimated time-ordered associations between KAKEN resources and Associate Professor promotion differ by gender?",
        "RQ3d. Among researchers who reach Associate Professor, do women and men show different post-promotion trajectories in KAKEN access, PI opportunities, funding, project category, and project-based collaboration networks?",
    ]:
        add_para(doc, rq)
    add_heading(doc, "1.3 Dissertation Structure", 2)
    add_para(doc, "The dissertation follows a deliberate division of labour. Study 1 identifies timing. Study 2 examines mobility pathways and organizational movement. Study 3 examines access to academic resources, project-based relational positions, and their time-ordered associations with promotion. This sequence prevents a single design from being asked to answer incompatible questions and keeps completed-transition evidence distinct from risk-set and sequence evidence.")
    add_picture(doc, figs / "framework.png", 6.8)
    add_caption(doc, "Figure 1. Dissertation logic: timing → mobility pathways → academic resources and returns.")


def add_literature(doc: Document) -> None:
    add_heading(doc, "2. Literature Review and Theoretical Framework", 1)
    add_para(doc, "The framework connects observed career histories to sociological accounts of organizations, mobility, status, linked lives, and cumulative advantage. Each theoretical component motivates an observable comparison or is explicitly treated as an unobserved mechanism.")
    add_heading(doc, "2.1 Gender Inequality in Japanese Academic Careers", 2)
    add_para(doc, "Gender inequality in Japanese academia appears in rank composition, advancement, research engagement, and leadership. Duration analysis in economics identified gender differences in promotion, studies of academic medicine documented rank disparities, and recent large-scale evidence shows pronounced heterogeneity across cohorts, fields, and stages (Takahashi & Takahashi, 2015; Nagano et al., 2022; Pei et al., 2026). Cross-national career evidence likewise emphasizes variation by country and discipline (Huang et al., 2020). These studies motivate separate analyses of transition timing, entry into promotion, and organizational pathways rather than a single comparison of final rank.")
    add_para(doc, "Care responsibilities and work-family organization are theoretically relevant because career opportunities may require travel, relocation, or periods of employment uncertainty. Research visits among Japanese bioscientists illustrate the career relevance of international mobility (Lawson & Shibayama, 2015). The present data do not observe marriage negotiations, childcare allocations, or private relocation costs. The proposal therefore treats these as plausible mechanisms to be evaluated cautiously, not as measured explanations. This distinction motivates gender-stratified descriptions and interaction models while preventing individual preferences from being inferred from observed mobility.")
    add_heading(doc, "2.2 Meritocracy, Gendered Evaluation, and Career-Relevant Academic Resources", 2)
    add_para(doc, "Academic evaluation combines formal criteria with institutional signals and relational judgments. Nielsen (2016) shows that formal recruitment and promotion procedures may retain gendered evaluation, while Moher et al. (2018) and Schimanski and Alperin (2018) caution that narrow performance indicators do not exhaust scholarly contribution. This dissertation therefore treats rank, institutional prestige, funding, and network position as career-relevant academic resources and signals, not as pure measures of merit.")
    add_para(doc, "The empirical implication is to separate access from intensity and status from performance. Study 2 measures whether researchers can enter and move among organizational locations. Study 3 distinguishes KAKEN participation from conditional funding amount and project role. Neither study interprets residual gender differences as direct proof of discrimination; they are conditional associations within measured data.")
    add_heading(doc, "2.3 Career Mobility, Status Attainment, and Segmented Academic Labour Markets", 2)
    add_para(doc, "Institutional mobility can alter job match, resources, colleagues, and status location. Allison and Long (1987) showed that interuniversity mobility is patterned by scientific productivity and institutional location. Fernández-Zubieta et al. (2016) found that mobility per se did not uniformly improve productivity; direction toward stronger departments mattered. These findings motivate measures of origin, destination, and direction rather than treating every institutional change as equivalent.")
    add_para(doc, "A segmented-labour-market perspective further implies that moves across national, public, private, research-institute, and nonacademic sectors may offer different opportunity structures. Study 2 will therefore classify institution type when the institutional property table becomes organized and will retain unresolved classifications rather than force them into a homogeneous university category. The relevant tests compare move incidence, direction, and promotion association across segments.")
    add_heading(doc, "2.4 Gendered Organizations, Linked Lives, and Spatial Constraints", 2)
    add_para(doc, "Gendered-organization theory challenges the image of the disembodied ideal worker. Organizational arrangements may reward continuous availability, geographic flexibility, and the capacity to prioritize career demands (Acker, 1990). In academic careers, institutional mobility can therefore function as a resource while the ability to move is itself unequally distributed. The central theoretical statement is: the ability to move may itself be an unequally distributed career resource.")
    add_para(doc, "Linked-lives and household-migration research treats mobility as coordinated across partners and dependants rather than as an isolated preference. Bielby and Bielby (1992) showed that family migration decisions are gendered, while Cañibano et al. (2016) found gender differences in the frequency, timing, duration, and distance of temporary research mobility. Study 2 consequently uses structural language about dual-career coordination, care obligations, geographic attachment, and unequal capacity to absorb relocation costs. These costs are not directly observed; geography will enter only after a validated institution crosswalk exists.")
    add_heading(doc, "2.5 Institutional Prestige, Organizational Scale, and Academic Mobility", 2)
    add_para(doc, "Academic labour markets are hierarchical. Hiring-network studies demonstrate strong prestige ordering and unequal placement flows (Clauset et al., 2015; Wapman et al., 2022). Study 2 operationalizes status movement using annual within-ranking percentiles calculated separately for THE, QS, and ARWU. It does not average raw ranks across systems. Three-year pre-move medians reduce annual noise, and broad tiers prevent minor fluctuations from being treated as substantive transitions.")
    add_para(doc, "Institution size captures a different organizational dimension. Student enrolment and full-time faculty indicate scale, graduate-training capacity, and potential teaching or administrative complexity, not quality. Annual percentiles and lagged medians improve comparability over time. Research institutes are retained as a distinct state rather than coded as zero-student universities. Prestige and size measures answer whether researchers move toward higher-status or larger research-training environments and whether those associations differ by gender.")
    add_heading(doc, "2.6 KAKEN Resources and Project Roles as Career-Relevant Academic Resources", 2)
    add_para(doc, "KAKEN is a central competitive funding system in Japan, with project categories that differ in purpose, scale, duration, and eligibility (JSPS, 2026a). Research on grant review reports mixed and context-dependent evidence regarding gendered evaluation, reinforcing the need to separate access, application context, and funding intensity (Sato et al., 2021). Project category is therefore retained from v_jp_researchers_kaken_project.research_category rather than treated as unavailable or collapsed at extraction. PI and Co-I are project-year roles, not permanent researcher attributes. Annual and cumulative role measures capture participation and leadership progression.")
    add_para(doc, "Funding access and funding intensity are analytically distinct. A two-part structure first models observed KAKEN participation and then examines positive direct funding among participants. Full-project, equal-fractional, and role-weighted exposures are compared. Every identifiable project member remains in fractional denominators even when not linked to the final researchmap cohort. This rule prevents the analytic sample from artificially inflating linked researchers' budget shares.")
    add_heading(doc, "2.7 Project-Based Collaboration Networks and Academic Careers", 2)
    add_para(doc, "Formal project participation creates a bipartite network between researchers and KAKEN projects. Its researcher projection measures project-based collaboration, not co-authorship, mentoring, or informal ties. Degree captures distinct formal collaborators, weighted degree repeated or multiple project ties, PI outdegree leadership reach through PI-led projects, and PI indegree inclusion in projects led by others. Early collaboration has been associated with later career outcomes in other settings, but those findings do not establish the same process in KAKEN data (Li et al., 2019).")
    add_para(doc, "The primary Study 3 model uses degree, PI outdegree, and PI indegree to limit collinearity. Weighted degree replaces degree in robustness models rather than entering simultaneously. Rolling windows are preferred when annual networks are sparse. These measures answer whether relational positions observed before promotion differ by gender and are associated with subsequent promotion risk.")
    add_heading(doc, "2.8 Cumulative Advantage and Post-Promotion Trajectories", 2)
    add_para(doc, "Cumulative-advantage theory describes how prior resources and recognition may be associated with later opportunities (Merton, 1968; DiPrete & Eirich, 2006). Funding research provides evidence that early success can be linked to later funding and continued system participation (Bol et al., 2018; Wang et al., 2019). These studies motivate cumulative prior-resource measures and the post-Associate-Professor extension, but they do not justify treating AP promotion as randomly assigned.")
    add_para(doc, "Study 3 therefore reports time-ordered resource-promotion associations and descriptive post-promotion trajectories. Event time aligns researchers around first observed AP appointment, while cohort-specific coverage and pre-event trends are reported. The analysis will not label a post-AP increase a causal promotion effect.")
    add_heading(doc, "2.9 Longitudinal, Event-History, and Sequence Approaches to Academic Careers", 2)
    add_para(doc, "Discrete-time event-history analysis retains researchers who have not yet experienced promotion and models annual transition risk. Flexible time-since-PhD functions avoid imposing a linear baseline hazard. Complementary log-log models are primary, logit models provide comparison, and Cox models serve as robustness checks. Predictor windows end at t−1 to prevent promotion-year information from entering the prediction of promotion in t.")
    add_para(doc, "Sequence analysis addresses a different question: how complete combinations of rank and institutional states unfold. Optimal Matching has an established sociological use for comparing career trajectories (Abbott & Tsay, 2000). Multichannel sequence analysis computes channel-specific distances before transparent combination, avoiding an unmanageably large compound state (Gauthier et al., 2010). Optimal Matching, transition-rate costs, and Dynamic Hamming distances will be compared; PAM is primary and Ward clustering a comparison. Gender is excluded from clustering and examined only after pathway types are established.")


def add_study1(doc: Document, media: Path) -> None:
    add_heading(doc, "3. Study 1: Gendered Promotion Timelines in Japanese Academia", 1)
    add_para(doc, "Study 1 is complete and provides the diagnostic starting point. It examines promotion timeline from PhD to Associate Professor and from Associate Professor to Full Professor using researchmap career trajectories. The published study uses 74,344 discipline-mapped PhD holders and explicitly distinguishes cohort composition, observed transition realization, and duration among completed transitions (Pei et al., 2026). The proposal does not recalculate the official sample.")
    add_heading(doc, "3.1 Purpose", 2)
    add_para(doc, "The study asks where gender differences appear across promotion stages, cohorts, and disciplines. Promotion timeline is the elapsed duration between validated milestones. It is not the same as the probability of entering or completing a transition.")
    add_heading(doc, "3.2 Data and Measurement", 2)
    add_para(doc, "The official discipline-mapped PhD cohort contains 74,344 researchers. The Stage 1 observed-transition sample contains 20,808 researchers and the Stage 2 observed-transition sample 10,747; positive-duration samples contain 18,794 and 10,416, respectively. Survival-based robustness checks address right-censoring. These samples are preserved as reported in the completed study.")
    add_picture(doc, media / "image3.png", 6.6)
    add_caption(doc, "Figure 2. Completed Study 1 official analytic and observed-transition samples. Source: Pei et al. (2026).")
    add_heading(doc, "3.3 Findings and Role in the Dissertation", 2)
    add_para(doc, "The completed analysis identifies the PhD-to-Associate-Professor transition as the principal site of observed heterogeneity. In positive-duration completed transitions, mean timelines are 7.83 years for women and 8.77 years for men from PhD to AP, and 8.22 years for women and 8.60 years for men from AP to FP. These conditional means must not be interpreted as evidence that women have a higher transition probability.")
    add_picture(doc, media / "image4.png", 6.6)
    add_caption(doc, "Figure 3. Completed Study 1 mean promotion timelines by gender and career stage. Source: Pei et al. (2026).")
    add_para(doc, "Study 1 motivates two next questions. Study 2 asks whether women and men reach similar destinations through different patterns of institutional stability and movement. Study 3 asks whether access to KAKEN resources and project-based relational positions differs before promotion and whether those resources are associated with first AP transition. Retaining right-censored researchers in Studies 2 and 3 extends, but does not eliminate, the selection limitations of completed-transition comparisons.")


def add_study2(doc: Document, figs: Path) -> None:
    add_heading(doc, "4. Study 2: Institutional Mobility as a Gendered Career Strategy", 1)
    add_heading(doc, "4.1 Research Objective", 2)
    add_para(doc, "Study 2 examines institutional mobility as both an organizational pathway and a potentially costly career resource. It asks who moves, when, in what direction, and with what observed career outcomes. The study does not assume that women are less willing to move. It evaluates gender differences in observed mobility while interpreting dual-career coordination, care obligations, and relocation costs as plausible but unmeasured structural mechanisms.")
    add_heading(doc, "4.2 Theoretical Expectations", 2)
    add_para(doc, "Moves to higher-prestige or more research-intensive institutions may be positively associated with promotion because they change organizational location, resources, and evaluation context. Lateral and downward moves may reflect strategic matching, temporary adjustment, constrained choice, or sector change; their meaning cannot be inferred from direction alone. The same observed move may carry different private costs and different career associations for women and men. These expectations are associational and will be tested through pre-specified interactions rather than universal claims.")
    add_heading(doc, "4.3 Sample and Career-History Scope", 2)
    add_para(doc, "The analytical population includes researchers with a valid identifier, validated PhD year, gender coded as woman or man for the primary comparison, field or discipline information, and sufficient career-history coverage. Researchers enter the AP risk set in the first fiscal year after PhD completion and leave at first observed AP appointment or the observation endpoint. Researchers who never reach AP remain as right-censored cases. Already holding AP at first observation, unresolved chronology, unreliable identity linkage, and insufficient coverage are separate exclusion reasons.")
    add_para(doc, "Career spells are expanded to fiscal-year records after resolving overlaps and selecting a canonical primary institution. The hierarchy is full-time academic appointment, highest rank, longest duration, principal listed affiliation, and deterministic tie-breaking. Multiple concurrent affiliations remain flagged. Missing history is NOT_OBSERVED, not academic exit.")
    add_heading(doc, "4.4 Institutional and Mobility Variables", 2)
    add_para(doc, "An institutional move occurs when the canonical primary institution changes between consecutive observed spells or researcher-years. The move table records origin_institution_id, destination_institution_id, move_year, institution_change, move_order, and career_year_at_move. Within-institution rank change is not an institutional move. Promotion is classified as INTERNAL_PROMOTION, EXTERNAL_PROMOTION, or PROMOTION_WITH_UNRESOLVED_MOVE_STATUS where timing permits.")
    add_para(doc, "Mobility frequency is represented by total moves, moves before and after AP, moves per observed year, early-career moves, and time since last move. Stability is measured separately through longest institutional spell, modal-institution share, institutional entropy, time to first move, and pre/post-AP stability. This distinction prevents one move count from conflating repeated short spells with one long stable appointment.")
    add_picture(doc, figs / "mobility_pipeline.png", 6.8)
    add_caption(doc, "Figure 4. Study 2 institutional-mobility construction and analytical outputs.")
    add_heading(doc, "4.5 Institutional Prestige and Size Measures", 2)
    add_para(doc, "THE, QS, and ARWU rankings remain separate. For institution j in year t, prestige_percentile_jt = (N_t − rank_jt)/(N_t − 1), with higher values indicating greater relative prestige. Origin and destination measures use the median percentile from t−3 to t−1 where available. A move is upward when destination minus origin is at least 10 percentile points, lateral when the absolute difference is below 10 points, and downward when the difference is at most −10 points. Five- and fifteen-point thresholds are robustness checks. Ranking-system availability and coverage are explicit fields; ranks are not projected backward before a system exists.")
    add_para(doc, "Institution-year size measures include undergraduate, master's, doctoral, professional-degree, graduate-total, and regular-total students, full-time faculty, and student-faculty ratio. Log counts, annual percentiles, and broad size tiers support comparison across years. Move-level changes use t−3 to t−1 medians. Doctoral enrolment approximates research-training capacity, not quality. Research institutes are not assigned zero enrolment.")
    add_para(doc, "Geographic measures are deferred until a validated institution-geography crosswalk exists. Institution-name strings will not be used to infer prefecture or distance. Once validated, prefecture change, region change, overseas move, great-circle distance, metropolitan transition, and return from overseas can be added. The Japanese practice of tanshin funin provides institutional context but is not treated as equivalent to observed academic mobility.")
    add_heading(doc, "4.6 Event-History Design", 2)
    add_para(doc, "The unit is researcher-year and the primary outcome is first observed AP transition in year t. Time-varying predictors include cumulative move count through t−1, move in t−1, time since last move, prior upward/lateral/downward history, lagged prestige and size change, institutional stability, and selected gender interactions. Complementary log-log is primary, logit is the comparison, and Cox models are robustness checks. Baseline hazard is modeled flexibly using time-since-PhD categories or restricted cubic splines.")
    add_para(doc, "All mobility predictors must precede year t. No institution change first visible in the promotion year is used as a pre-promotion predictor unless the ordering of move and promotion can be established. Models report exponentiated estimates, predicted probabilities, average marginal effects, and clustered uncertainty. Alternative samples exclude ambiguous AP years, unresolved institutions, and low-coverage careers.")
    add_heading(doc, "4.7 Multichannel Sequence Analysis", 2)
    add_para(doc, "Career-year sequences are aligned from PhD, primarily over years 0–15, with 0–10 and 0–20 sensitivity windows. The rank channel records training, postdoctoral/nonfaculty, assistant-professor-or-lecturer, AP, FP, other academic, nonacademic, and NOT_OBSERVED states. The prestige channel uses TIER_1 through TIER_4, UNRANKED, RESEARCH_INSTITUTE, OVERSEAS_UNCLASSIFIED, NONACADEMIC, and NOT_OBSERVED. The mobility channel records same institution, internal role change, within-type move, across-type move, research-institute move, overseas move, return, nonacademic move, and NOT_OBSERVED.")
    add_para(doc, "Separate standardized distance matrices are computed for each channel using constant-cost Optimal Matching, transition-rate costs, and Dynamic Hamming where appropriate. Equal channel weights are primary. PAM/k-medoids is the primary clustering method and Ward hierarchical clustering a comparison. Candidate solutions are judged by silhouette, size balance, within-cluster heterogeneity, resampling stability, agreement across distances, and substantive interpretability. Gender is excluded from clustering; clusters are named from medoids and state distributions before gender composition is inspected.")
    add_heading(doc, "4.8 Equal-Destination Comparison", 2)
    add_para(doc, "Among researchers reaching the same final rank, especially AP, the analysis compares move count, upward-move frequency, total prestige gain, institution-size change, stability, pathway membership, and promotion timing. This design tests whether equal destinations conceal unequal observed pathways. It remains conditional on reaching the destination and does not replace the full risk-set analysis.")
    add_heading(doc, "4.9 Interpretation and Limitations", 2)
    add_para(doc, "Study 2 estimates mobility-career associations, not returns identified by exogenous mobility. Mobility is selected by unobserved ability, opportunity, family circumstances, recruitment, and career preferences. Prestige rankings omit many institutions and change coverage over time; institution type is not yet fully organized; geography is unavailable pending validation; and family costs are theoretically relevant but not observed. Results will therefore distinguish observed mobility patterns, hypothesized mechanisms, and unobserved private costs.")


def add_study3(doc: Document, figs: Path) -> None:
    add_heading(doc, "5. Study 3: Gendered Access to and Returns from KAKEN Resources and Project-Based Collaboration Networks", 1)
    add_heading(doc, "5.1 Research Objective", 2)
    add_para(doc, "Study 3 examines gender differences in observed KAKEN access, project roles, funding intensity, original project category, and project-based network position. It then asks whether resources measured before year t are associated with first observed AP transition in t and whether selected associations differ by gender. The term returns denotes estimated associations with observed career outcomes, not causal effects.")
    add_heading(doc, "5.2 Promotion-Risk Population and Time Ordering", 2)
    add_para(doc, "The population follows the Study 2 AP risk-set rules but uses KAKEN and network predictors. Researchers with no observed KAKEN participation remain in the panel after linkage validation. LINKED_CONFIDENT, LINKED_PROBABLE, NO_OBSERVED_KAKEN, UNRESOLVED_LINKAGE, and SOURCE_COVERAGE_LIMITED are distinct states. Post-event years are excluded.")
    add_picture(doc, figs / "kaken_timing.png", 6.8)
    add_caption(doc, "Figure 5. Study 3 temporal ordering: KAKEN and network predictors end before promotion year t.")
    add_heading(doc, "5.3 KAKEN Access, Roles, Funding, and Project Category", 2)
    add_para(doc, "Annual measures distinguish any participation, PI, Co-I, other role, simultaneous PI and Co-I, project counts, cumulative prior PI/Co-I years, first PI year, and Co-I-to-PI transition. PI and Co-I remain project-specific roles. Project category comes directly from v_jp_researchers_kaken_project.research_category and is retained in its original form; simplified categories are presentation or robustness variables only.")
    add_para(doc, "Funding uses direct project-year budget records. Exposure assigns the full project amount to each participant; equal fractional exposure divides by all identifiable members; role-weighted exposure uses documented PI, Co-I, and collaborator weights while retaining unlinked members in the denominator. The main two-part specification separates KAKEN access from log_positive_direct_funding among positive observations. Field-year standardized positive funding adjusts for field and calendar differences without being interpreted as merit.")
    add_heading(doc, "5.4 KAKEN-Based Collaboration Networks", 2)
    add_para(doc, "For each project-year, all identifiable member pairs form undirected edge events; PI-to-Co-I edges form the directed leadership network. Repeated collaborations across projects remain distinct before researcher-pair-year aggregation. Annual and rolling networks produce degree, weighted degree, PI outdegree, and PI indegree. The primary model uses degree and directed measures; weighted degree substitutes for degree in robustness analysis. Betweenness and eigenvector centrality remain optional because sparse components and computational availability may reduce reliability.")
    add_heading(doc, "5.5 Gender Differences in Resource Access and Associations", 2)
    add_para(doc, "Tier 1 describes KAKEN access, roles, funding, category, and network position by gender, field, cohort, career year, and institution where valid. Tier 2 estimates resource-promotion associations with complementary log-log models controlling for flexible baseline hazard, field, cohort, calendar year, and validated institutional context. Tier 3 adds a limited set of pre-specified gender interactions for access, PI status, positive funding, degree, and PI outdegree. Interactions are interpreted through predicted probabilities and marginal effects.")
    add_para(doc, "Primary timing is t−1 and past3, defined as t−3 through t−1. Past4, past5, and cumulative prior exposure are robustness specifications. Raw, logged positive, and field-year standardized funding enter separate model families. Degree and weighted degree do not enter the same main specification. High-linkage-quality and high-career-coverage samples assess sensitivity to missingness.")
    add_heading(doc, "5.6 Post-Associate-Professor Resource Trajectories", 2)
    add_para(doc, "RQ3d is a secondary extension among researchers with validated AP year. Event time equals fiscal year minus AP year, with a proposed −5 to +5 window. Outcomes include KAKEN access, PI status, project count, positive direct funding, project category, degree, PI outdegree, and PI indegree. Cohort-specific support and observation coverage are reported. The design describes pre/post trajectories and gender differences; it does not estimate a causal effect of AP promotion.")
    add_heading(doc, "5.7 Interpretation and Limitations", 2)
    add_para(doc, "Missing KAKEN records may reflect non-participation, linkage failure, or source coverage. Project networks omit co-authorship, mentoring, and informal collaboration. Budget exposure is not the same as money personally controlled, and category meanings vary historically. Lagging improves temporal ordering but cannot remove confounding by research productivity, institutional selection, family circumstances, or unobserved ability. Conclusions will use the terms observed access, promotion risk, time-ordered association, and post-promotion resource trajectory.")


def add_methods(doc: Document) -> None:
    add_heading(doc, "6. Data and Methodological Strategy", 1)
    add_heading(doc, "6.1 Data Sources and Shared Foundations", 2)
    add_para(doc, "The shared researcher spine links validated researchmap identifiers, doctoral-degree timing, gender, discipline, and career-history coverage. researchmap is operated by JST and provides researcher-entered or institution-supplied profiles and achievements (JST, 2026). KAKEN project membership, project metadata, annual budgets, and original research_category are linked through documented project and researcher identifiers. Database credentials are environment variables and never enter scripts or outputs.")
    add_para(doc, "Shared intermediate data consist of a career-year panel, institution-year panel, institutional-move table, prestige-year table, institution student-year table, researcher-project-year panel, KAKEN resource-year panel, and KAKEN network-year panel. Each output carries provenance, row counts, missingness, duplicate diagnostics, and source-coverage flags.")
    add_heading(doc, "6.2 Identity, Rank, and Institution Harmonization", 2)
    add_para(doc, "Researcher linkage distinguishes the KAKEN website user identifier from the JPR researchmap userId; the dissertation cohort uses the JPR userId. Unlinked project members remain in team-size and fractional-budget denominators. Rank harmonization extracts first AP from research-experience titles including 准教授 and associate professor; assistant-professor views do not define the outcome. A title-audit table records raw title, language, standardized rank, frequency, ambiguity, and manual-review status.")
    add_para(doc, "Institution names are canonicalized through identifiers and audited aliases. Institution type and official properties will replace temporary string heuristics once jp_researchers_institute_property is available. A comparison report will document reclassification. No city or prefecture is inferred from free text.")
    add_heading(doc, "6.3 Derived Study Panels", 2)
    add_para(doc, "Study 2 derives a mobility risk panel and three principal sequence channels from the career and institution foundations. Study 3 derives a KAKEN promotion-risk panel and a post-AP event-time panel from the resource and network foundations. AP timing and researcher spine are shared, but Study 2 mobility sequences are never used as predictors in Study 3 if they include post-promotion information.")
    add_heading(doc, "6.4 Missingness, Coverage, and Reproducibility", 2)
    add_para(doc, "Explicit flags cover missing PhD year, gender, unresolved rank, ambiguous AP year, unmatched researcher, unmatched project, missing role, missing budget, incomplete institution history, ranking coverage, and insufficient observation. Missingness is reported by gender, field, cohort, and institution type. NO_OBSERVED_KAKEN and UNRESOLVED_LINKAGE never share a code.")
    add_para(doc, "All extraction and transformation decisions are version controlled. Run manifests contain source tables, extraction timestamp, configuration, command, output schema, row counts, checksums, and validation results. Data too large or confidential for Git are represented by manifests rather than embedded credentials or raw extracts.")
    add_heading(doc, "6.5 Ethics and Interpretation Boundaries", 2)
    add_para(doc, "The analysis uses career and project records for population-level research. Outputs will minimize disclosure, avoid publishing unnecessary identifiable records, and report aggregate results. Gender coding is limited to the available binary classification for the primary comparison and does not represent the full range of gender identities. Observational designs cannot identify causal effects of mobility, funding, project roles, or promotion. Unmeasured family circumstances and private costs are not attributed to individuals.")


def add_contributions(doc: Document) -> None:
    add_heading(doc, "7. Expected Contributions and Work Plan", 1)
    add_heading(doc, "7.1 Expected Contributions", 2)
    add_para(doc, "The dissertation offers four contributions. First, Study 1 establishes where gender differences appear in observed promotion timing and why completed-transition comparisons require censoring and selection caveats. Second, Study 2 contributes a sociological account of mobility as a gendered career strategy by linking organizational boundaries, status direction, scale, stability, and promotion risk. Third, Study 3 contributes to sociology of science and science of science by distinguishing access to KAKEN from funding intensity and by measuring formal project-based relational positions. Fourth, the integrated data architecture demonstrates how event-history and multichannel sequence methods can be combined without leaking future information into pre-promotion models.")
    add_heading(doc, "7.2 Work Plan", 2)
    rows = [
        ["Period", "Main task", "Output/milestone"],
        ["July 2026", "Finalize proposal revision; verify literature, terminology, and citations; preserve Study 1.", "Tracked and clean proposal; audit files."],
        ["August 2026", "Build institution crosswalk; organize institute properties; validate career spells and AP years.", "Institution and rank validation reports."],
        ["September 2026", "Acquire and validate THE/QS/ARWU and institution student-year data; construct move events.", "Prestige, size, and mobility panels."],
        ["October 2026", "Estimate Study 2 risk models; compute sequence distances and clusters; run equal-destination analyses.", "Study 2 tables, figures, and diagnostics."],
        ["November 2026", "Refresh KAKEN panels using original research_category; validate roles, budgets, and linkage denominators.", "Study 3 resource and linkage reports."],
        ["December 2026", "Estimate Study 3 resource-access and promotion-risk models; build post-AP extension.", "Study 3 tables, figures, and diagnostics."],
        ["January–March 2027", "Integrate results, robustness checks, limitations, and chapter drafts.", "Full dissertation chapter drafts."],
    ]
    table = doc.add_table(rows=1, cols=3)
    for i, value in enumerate(rows[0]):
        table.rows[0].cells[i].text = value
    for row in rows[1:]:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
    format_table(table, widths=[1400, 4300, 3000], font_size=8)
    add_heading(doc, "7.3 Conclusion", 2)
    add_para(doc, "The revised dissertation follows the progression timing → mobility pathways → academic resources and returns. It extends completed-transition evidence by retaining researchers still at risk, identifies institutional movement as an organizational process rather than an individual preference, and locates KAKEN access and project networks within a time-ordered resource framework. The result is a coherent, auditable design for describing gendered academic careers in Japan while maintaining clear limits on causal interpretation.")


STUDY2_VARIABLES = [
    ["Institutional move", "institution_change", "Canonical primary institution differs across consecutive observed spells/years", "At transition; lagged for risk model", "researchmap career history", "Separates organizational crossing from rank change", "Boundary-crossing mobility", "RQ2a–c", "Gaps may mimic moves", "Require adjacent observed years"],
    ["Move timing", "move_year; career_year_at_move", "Fiscal year and years since PhD at each move", "Observed event", "Career-year panel", "Locates mobility by career stage", "Timing of mobility", "RQ2a", "Spell dates may be coarse", "Alternative fiscal-year assignment"],
    ["Move frequency", "total_moves; moves_per_observed_year", "Count and exposure-adjusted rate", "Through t−1 or full sequence", "Move table", "Separates count from observation length", "Mobility intensity", "RQ2a,c,d", "Sensitive to coverage", "High-coverage sample"],
    ["Promotion route", "promotion_move_status", "Internal, external, or unresolved at first AP/FP", "Promotion year with ordering check", "Rank + move tables", "Distinguishes loyalty from boundary crossing", "Organizational route to advancement", "RQ2b,c,e", "Move/promotion order may be unknown", "Exclude unresolved"],
    ["Prestige", "prestige_percentile", "Within-system annual percentile, separately for THE/QS/ARWU", "t−3:t−1 median", "Ranking-year tables", "Adjusts for annual coverage", "Institutional status", "RQ2b,c,e", "Rankings omit institutions", "System-specific and tier models"],
    ["Move direction", "upward/lateral/downward", "Destination-origin percentile change using ±10-point rule", "Pre-move medians", "Prestige-year + move", "Avoids treating all moves alike", "Status mobility", "RQ2b,c,e", "Threshold-dependent", "±5 and ±15 points"],
    ["Prestige tier", "TIER_1–TIER_4 etc.", "Broad percentile/structural states", "Annual sequence", "Prestige-year + type", "Limits sequence state proliferation", "Status location", "RQ2d,e", "Cut points simplify hierarchy", "Alternative cut points"],
    ["Institution size", "log_total_students; percentiles", "Student/faculty counts, logs, and annual percentiles", "t−3:t−1 median", "Institution student-year", "Makes scale comparable over time", "Organizational scale", "RQ2b,c,e", "Scale is not quality", "Doctoral/graduate alternatives"],
    ["Research training", "log_doctoral_students", "Log doctoral enrolment with separate missing state", "Lagged median", "Institution student-year", "Targets graduate-training capacity", "Research-training environment", "RQ2b,e", "Field mix not observed", "Graduate-total measure"],
    ["Stability", "longest_spell; modal_share; entropy", "Duration, concentration, and diversity of institution history", "Through t−1/full sequence", "Career-year panel", "Separates stability from move count", "Institutional attachment/fragmentation", "RQ2a,c,d,e", "Affected by missing years", "High-coverage sample"],
    ["Geographic mobility", "prefecture/region/distance", "Validated origin-destination geography only", "At move", "Future crosswalk", "Institution change may not imply relocation", "Spatial mobility burden", "RQ2a,c", "Not currently available", "Domestic/overseas only"],
    ["AP transition", "AP_event", "First validated AP appointment in fiscal year t", "Year t outcome", "research_experience", "Defines annual promotion risk", "Career advancement", "RQ2b,c", "First observation may already be AP", "Exclude ambiguous AP years"],
]

STUDY3_VARIABLES = [
    ["KAKEN access", "KAKEN_access_t1", "Any linked project participation in t−1", "t−1; past3", "Researcher-project-year", "Separates system entry from intensity", "Observed resource access", "RQ3a–c", "No record may be linkage failure", "High-linkage sample"],
    ["Project role", "PI_t1; CoI_t1", "Project-year PI/Co-I indicators", "t−1; rolling/cumulative", "Member role metadata", "Roles vary by project and year", "Leadership and formal participation", "RQ3a–d", "Role strings may be ambiguous", "Strict/expanded harmonization"],
    ["Leadership progression", "first_PI_year; CoI_to_PI", "First PI and prior Co-I-to-PI transition", "Prior to t", "Researcher-project-year", "Captures progression beyond annual status", "Project leadership trajectory", "RQ3a,c,d", "Left censoring", "Cohort-restricted sample"],
    ["Project breadth", "project_count", "Distinct active projects, including PI and Co-I counts", "Annual; past3", "Researcher-project-year", "Captures concurrent breadth/workload", "Resource portfolio", "RQ3a–d", "Projects differ in scale", "Role-specific counts"],
    ["Direct funding", "direct_funding", "Sum of annual direct project budgets", "t−1; past3", "Project-year budget", "Direct amount is closer to usable resources", "Funding intensity", "RQ3a–d", "Exposure is not personal control", "Full vs fractional"],
    ["Positive funding", "log_positive_direct_funding", "log(positive direct funding) among participants", "Lagged", "Resource-year panel", "Avoids conflating zero with access", "Conditional intensity", "RQ3b,c", "Conditional sample", "Raw and standardized"],
    ["Fractional funding", "equal/role_weighted_fraction", "Budget divided using all identifiable members", "Annual; lagged", "Project member + budget", "Prevents linked-cohort denominator bias", "Allocated exposure", "RQ3a–c", "Weights are assumptions", "Equal vs role-weighted"],
    ["Project category", "research_category", "Original v_jp_researchers_kaken_project.research_category", "Project-year; lagged", "KAKEN project view", "Retains scheme heterogeneity", "Funding-program context", "RQ3a–d", "Categories change historically", "Simplified groups"],
    ["Degree", "degree", "Distinct project collaborators in window", "Annual; past3", "Undirected projection", "Counts formal collaborator reach", "Project relational breadth", "RQ3a–d", "Not co-authorship", "Weighted degree replacement"],
    ["Weighted degree", "weighted_degree", "Repeated/shared-project tie weight", "Annual; rolling", "Pair-year network", "Captures repeated collaboration", "Tie intensity", "RQ3a,d", "Correlated with degree", "Use instead of degree"],
    ["PI outdegree", "PI_outdegree", "Distinct Co-Is reached by PI→Co-I ties", "Annual; past3", "Directed network", "Measures leadership reach", "Directed project leadership", "RQ3a–d", "Depends on role quality", "Strict-role network"],
    ["PI indegree", "PI_indegree", "Distinct PIs linking to researcher as Co-I", "Annual; past3", "Directed network", "Measures inclusion by leaders", "Directed inclusion", "RQ3a–d", "Not mentoring", "Strict-role network"],
    ["AP transition", "AP_event", "First validated AP appointment in t", "Outcome t", "research_experience", "Defines promotion risk", "Career advancement", "RQ3b,c", "Rank ambiguity", "Validated-only sample"],
    ["Post-AP time", "event_time", "Fiscal year minus validated AP year", "−5 to +5", "AP + resource panels", "Aligns descriptive trajectories", "Post-promotion accumulation", "RQ3d", "No causal treatment", "Alternative windows"],
]


def add_variable_table(doc: Document, title: str, rows: list[list[str]]) -> None:
    add_heading(doc, title, 2)
    headers = ["Construct", "Variable", "Operational definition", "Timing", "Source",
               "Why constructed this way", "Conceptual meaning", "Research question answered",
               "Main limitation", "Robustness version"]
    table = doc.add_table(rows=1, cols=len(headers))
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
    widths = [900, 1050, 1500, 950, 1100, 1350, 1150, 1200, 1200, 1100]
    format_table(table, widths=widths, font_size=6.5)


def add_appendix(doc: Document) -> None:
    landscape = doc.add_section(WD_SECTION.NEW_PAGE)
    landscape.orientation = WD_ORIENT.LANDSCAPE
    landscape.page_width, landscape.page_height = landscape.page_height, landscape.page_width
    landscape.left_margin = Inches(0.45)
    landscape.right_margin = Inches(0.45)
    landscape.top_margin = Inches(0.55)
    landscape.bottom_margin = Inches(0.55)
    add_heading(doc, "Appendix A. Variable Construction and Validation", 1)
    add_heading(doc, "A.1 Source and Data-Status Codes", 2)
    add_para(doc, "Core status codes are LINKED_CONFIDENT, LINKED_PROBABLE, NO_OBSERVED_KAKEN, UNRESOLVED_LINKAGE, SOURCE_COVERAGE_LIMITED, RANK_VALIDATED, RANK_AMBIGUOUS, INSTITUTION_VALIDATED, INSTITUTION_UNRESOLVED, OBSERVED_ZERO, NOT_OBSERVED, and RIGHT_CENSORED. Codes distinguish absence, non-participation, linkage failure, and source limitation.")
    add_variable_table(doc, "A.2 Study 2 Mobility, Prestige, Size, and Sequence Variables", STUDY2_VARIABLES)
    add_variable_table(doc, "A.3 Study 3 KAKEN and Network Variables", STUDY3_VARIABLES)
    portrait = doc.add_section(WD_SECTION.NEW_PAGE)
    portrait.orientation = WD_ORIENT.PORTRAIT
    portrait.page_width, portrait.page_height = portrait.page_height, portrait.page_width
    portrait.left_margin = Inches(0.9)
    portrait.right_margin = Inches(0.9)
    portrait.top_margin = Inches(0.8)
    portrait.bottom_margin = Inches(0.8)
    add_heading(doc, "A.4 Required Diagnostics", 2)
    add_bullets(doc, [
        "Sample inclusion and exclusion counts, event count, censoring rate, risk-set size by year, and observation coverage.",
        "Missingness and linkage quality by gender, field, cohort, and institution type.",
        "Rank-title, institution-alias, project-role, project-category, budget, and ranking-coverage audits.",
        "Sparse cells, multicollinearity, interaction support, influential observations, calibration, and discrimination.",
        "Sequence missingness, NOT_OBSERVED dominance, cluster stability, silhouette, medoids, and high-coverage replication.",
        "Temporal leakage checks proving that promotion-year and later KAKEN/network information do not enter year-t risk models.",
    ])


REFERENCES = [
    "Abbott, A., & Tsay, A. (2000). Sequence analysis and optimal matching methods in sociology: Review and prospect. Sociological Methods & Research, 29(1), 3–33. https://doi.org/10.1177/0049124100029001001",
    "Acker, J. (1990). Hierarchies, jobs, bodies: A theory of gendered organizations. Gender & Society, 4(2), 139–158. https://doi.org/10.1177/089124390004002002",
    "Allison, P. D., & Long, J. S. (1987). Interuniversity mobility of academic scientists. American Sociological Review, 52(5), 643–652. https://doi.org/10.2307/2095600",
    "Bielby, W. T., & Bielby, D. D. (1992). I will follow him: Family ties, gender-role beliefs, and reluctance to relocate for a better job. American Journal of Sociology, 97(5), 1241–1267. https://doi.org/10.1086/229901",
    "Bol, T., de Vaan, M., & van de Rijt, A. (2018). The Matthew effect in science funding. Proceedings of the National Academy of Sciences, 115(19), 4887–4890. https://doi.org/10.1073/pnas.1719557115",
    "Cañibano, C., Fox, M. F., & Otamendi, F. J. (2016). Gender and patterns of temporary mobility among researchers. Science and Public Policy, 43(3), 320–331. https://doi.org/10.1093/scipol/scv042",
    "Clauset, A., Arbesman, S., & Larremore, D. B. (2015). Systematic inequality and hierarchy in faculty hiring networks. Science Advances, 1(1), e1400005. https://doi.org/10.1126/sciadv.1400005",
    "DiPrete, T. A., & Eirich, G. M. (2006). Cumulative advantage as a mechanism for inequality: A review of theoretical and empirical developments. Annual Review of Sociology, 32, 271–297. https://doi.org/10.1146/annurev.soc.32.061604.123127",
    "Fernández-Zubieta, A., Geuna, A., & Lawson, C. (2016). Productivity pay-offs from academic mobility: Should I stay or should I go? Industrial and Corporate Change, 25(1), 91–114. https://doi.org/10.1093/icc/dtv034",
    "Gauthier, J.-A., Widmer, E. D., Bucher, P., & Notredame, C. (2010). Multichannel sequence analysis applied to social science data. Sociological Methodology, 40(1), 1–38. https://doi.org/10.1111/j.1467-9531.2010.01227.x",
    "Huang, J., Gates, A. J., Sinatra, R., & Barabási, A.-L. (2020). Historical comparison of gender inequality in scientific careers across countries and disciplines. Proceedings of the National Academy of Sciences, 117(9), 4609–4616. https://doi.org/10.1073/pnas.1914221117",
    "Japan Science and Technology Agency. (2026). About the researchmap project. https://researchmap.jp/public/about/operations",
    "Japan Society for the Promotion of Science. (2026a). KAKENHI: Types of grants programs. https://www.jsps.go.jp/english/e-grants/grants01.html",
    "Japan Society for the Promotion of Science. (2026b). KAKENHI award trends. https://www.jsps.go.jp/english/e-grants/award_trends.html",
    "Lawson, C., & Shibayama, S. (2015). International research visits and careers: An analysis of bioscience academics in Japan. Science and Public Policy, 42(5), 690–710. https://doi.org/10.1093/scipol/scu084",
    "Li, W., Aste, T., Caccioli, F., & Livan, G. (2019). Early coauthorship with top scientists predicts success in academic careers. Nature Communications, 10, 5170. https://doi.org/10.1038/s41467-019-13130-4",
    "Merton, R. K. (1968). The Matthew effect in science. Science, 159(3810), 56–63. https://doi.org/10.1126/science.159.3810.56",
    "Moher, D., Naudet, F., Cristea, I. A., Miedema, F., Ioannidis, J. P. A., & Goodman, S. N. (2018). Assessing scientists for hiring, promotion, and tenure. PLOS Biology, 16(3), e2004089. https://doi.org/10.1371/journal.pbio.2004089",
    "Nagano, N., Watari, T., Tamaki, Y., & Onigata, K. (2022). Japan's academic barriers to gender equality as seen in a comparison of public and private medical schools. Women's Health Reports, 3(1), 115–123. https://doi.org/10.1089/whr.2021.0082",
    "Nielsen, M. W. (2016). Limits to meritocracy? Gender in academic recruitment and promotion processes. Science and Public Policy, 43(3), 386–399. https://doi.org/10.1093/scipol/scv052",
    "Pei, R., Lyu, Z., Wang, G., Wang, Z., Ye, M., & Fan, X. (2026). Gender inequality in academic promotion trajectories in Japan. Scientific Reports. https://doi.org/10.1038/s41598-026-54562-5",
    "Sato, S., Gygax, P. M., Randall, J., & Schmid Mast, M. (2021). The leaky pipeline in research grant peer review and funding decisions: Challenges and future directions. Higher Education, 82, 145–162. https://doi.org/10.1007/s10734-020-00626-y",
    "Schimanski, L. A., & Alperin, J. P. (2018). The evaluation of scholarship in academic promotion and tenure processes: Past, present, and future. F1000Research, 7, 1605. https://doi.org/10.12688/f1000research.16493.1",
    "Statistics Bureau of Japan. (2025). Statistical handbook of Japan 2025. https://www.stat.go.jp/english/data/handbook/",
    "Takahashi, A. M., & Takahashi, S. (2015). Gender promotion differences in economics departments in Japan: A duration analysis. Journal of Asian Economics, 41, 1–19. https://doi.org/10.1016/j.asieco.2015.09.002",
    "Wang, Y., Jones, B. F., & Wang, D. (2019). Early-career setback and future career impact. Nature Communications, 10, 4331. https://doi.org/10.1038/s41467-019-12189-3",
    "Wapman, K. H., Zhang, S., Clauset, A., & Larremore, D. B. (2022). Quantifying hierarchy and dynamics in US faculty hiring and retention. Nature, 610, 120–127. https://doi.org/10.1038/s41586-022-05222-x",
    "Watanabe, M. (2025). Gender inequality in international research engagement amid transformation to global and neoliberal academia: The case of Japan. Gender, Work & Organization, 32(5), 1812–1822. https://doi.org/10.1111/gwao.13224",
]


def add_references(doc: Document) -> None:
    add_heading(doc, "References", 1)
    for ref in REFERENCES:
        p = doc.add_paragraph(ref)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_together = True


def build_clean(source: Path, out: Path, work: Path) -> None:
    global CJK_IMAGE_PATH
    doc = Document(source)
    clear_body(doc)
    configure_styles(doc)
    media = work / "source_media"
    media.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(source) as source_zip:
        for member in source_zip.namelist():
            if member.startswith("word/media/") and not member.endswith("/"):
                (media / Path(member).name).write_bytes(source_zip.read(member))
    logo = media / "image1.png"
    if not logo.exists():
        raise FileNotFoundError(f"Expected title-page logo was not found in source DOCX: {logo}")
    figures = work / "figures"
    figures.mkdir(parents=True, exist_ok=True)
    make_framework_figure(figures / "framework.png")
    make_mobility_figure(figures / "mobility_pipeline.png")
    make_kaken_timing_figure(figures / "kaken_timing.png")
    CJK_IMAGE_PATH = figures / "junkyoju.png"
    make_cjk_title_pattern(CJK_IMAGE_PATH)
    add_title_page(doc, logo)
    add_front_matter(doc)
    add_intro(doc, figures)
    add_literature(doc)
    add_study1(doc, media)
    add_study2(doc, figures)
    add_study3(doc, figures)
    add_methods(doc)
    add_contributions(doc)
    add_appendix(doc)
    add_references(doc)
    out.parent.mkdir(parents=True, exist_ok=True)
    raw = out.with_suffix(".raw.docx")
    doc.save(raw)
    strip_comments(raw, out)
    raw.unlink()


def _xml_bytes(root: etree._Element) -> bytes:
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone="yes")


def strip_comments(source: Path, out: Path) -> None:
    """Remove comment anchors, relationships, content types, and comment parts."""
    omitted_prefixes = ("word/comments", "word/people.xml")
    with zipfile.ZipFile(source, "r") as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            name = info.filename
            if name.startswith(omitted_prefixes):
                continue
            data = zin.read(name)
            if name == "word/document.xml":
                root = etree.fromstring(data)
                for tag in ("commentRangeStart", "commentRangeEnd", "commentReference"):
                    for node in root.xpath(f".//w:{tag}", namespaces=NS):
                        node.getparent().remove(node)
                data = _xml_bytes(root)
            elif name == "word/_rels/document.xml.rels":
                root = etree.fromstring(data)
                for rel in list(root):
                    rel_type = rel.get("Type", "")
                    if "comments" in rel_type or rel_type.endswith("/people"):
                        root.remove(rel)
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone="yes")
            elif name == "[Content_Types].xml":
                root = etree.fromstring(data)
                for node in list(root):
                    part = node.get("PartName", "")
                    if part.startswith("/word/comments") or part == "/word/people.xml":
                        root.remove(node)
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone="yes")
            zout.writestr(info, data)


def make_tracked(clean: Path, out: Path, title_paragraphs: int = 13) -> None:
    """Mark revised body runs as true tracked insertions.

    The proposal is extensively restructured, so the tracked copy marks the full
    revised body after the preserved institutional title block as inserted text.
    Accepting changes yields the same substantive text as the clean copy.
    """
    when = dt.datetime.now(dt.timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")
    with zipfile.ZipFile(clean, "r") as zin:
        doc_root = etree.fromstring(zin.read("word/document.xml"))
        settings = etree.fromstring(zin.read("word/settings.xml"))
        if settings.find("w:trackRevisions", namespaces=NS) is None:
            settings.insert(0, etree.Element(f"{{{W_NS}}}trackRevisions"))
        body = doc_root.find("w:body", namespaces=NS)
        paragraphs = body.findall("w:p", namespaces=NS)
        first_body = paragraphs[min(title_paragraphs, len(paragraphs) - 1)] if paragraphs else None
        start_seen = False
        cid = 1000
        for child in list(body):
            if child is first_body:
                start_seen = True
            if not start_seen or child.tag == f"{{{W_NS}}}sectPr":
                continue
            for p in child.xpath(".//w:p", namespaces=NS) if child.tag != f"{{{W_NS}}}p" else [child]:
                for run in list(p):
                    if run.tag != f"{{{W_NS}}}r":
                        continue
                    idx = p.index(run)
                    ins = etree.Element(f"{{{W_NS}}}ins")
                    ins.set(f"{{{W_NS}}}id", str(cid))
                    ins.set(f"{{{W_NS}}}author", "Codex")
                    ins.set(f"{{{W_NS}}}date", when)
                    cid += 1
                    p.remove(run)
                    ins.append(run)
                    p.insert(idx, ins)
        overrides = {"word/document.xml": _xml_bytes(doc_root), "word/settings.xml": _xml_bytes(settings)}
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
            for info in zin.infolist():
                data = overrides.get(info.filename, zin.read(info.filename))
                zout.writestr(info, data)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", required=True, type=Path)
    parser.add_argument("--out-dir", required=True, type=Path)
    parser.add_argument("--work-dir", required=True, type=Path)
    args = parser.parse_args()
    source = args.source.resolve()
    out_dir = args.out_dir.resolve()
    work_dir = args.work_dir.resolve()
    out_dir.mkdir(parents=True, exist_ok=True)
    work_dir.mkdir(parents=True, exist_ok=True)
    clean = out_dir / "Doctoral_Dissertation_Proposal_PEI_Rongkang_revised_clean.docx"
    tracked = out_dir / "Doctoral_Dissertation_Proposal_PEI_Rongkang_revised_tracked.docx"
    build_clean(source, clean, work_dir)
    make_tracked(clean, tracked)
    print(json.dumps({"clean": str(clean), "tracked": str(tracked)}, indent=2))


if __name__ == "__main__":
    main()
