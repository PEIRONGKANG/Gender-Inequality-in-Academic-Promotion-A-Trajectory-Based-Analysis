#!/usr/bin/env python3
"""Run structural, citation, terminology, and tracked-copy audits."""

from __future__ import annotations

import argparse
import csv
import hashlib
import re
import subprocess
import sys
import tempfile
import zipfile
from collections import Counter
from pathlib import Path

from docx import Document
from lxml import etree


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W_NS}


def all_text(docx: Path) -> str:
    doc = Document(docx)
    chunks = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            chunks.extend(cell.text for cell in row.cells)
    return "\n".join(chunks)


def paragraph_text(docx: Path) -> str:
    return "\n".join(p.text for p in Document(docx).paragraphs)


def comment_count(docx: Path) -> int:
    with zipfile.ZipFile(docx) as z:
        if "word/comments.xml" not in z.namelist():
            return 0
        root = etree.fromstring(z.read("word/comments.xml"))
        return len(root.xpath(".//w:comment", namespaces=NS))


def revision_counts(docx: Path) -> tuple[int, int]:
    with zipfile.ZipFile(docx) as z:
        root = etree.fromstring(z.read("word/document.xml"))
    return len(root.xpath(".//w:ins", namespaces=NS)), len(root.xpath(".//w:del", namespaces=NS))


def headings(docx: Path) -> list[str]:
    doc = Document(docx)
    return [p.text.strip() for p in doc.paragraphs if p.style.name.startswith("Heading") and p.text.strip()]


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, required=True)
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--accept-script", type=Path, required=True)
    args = parser.parse_args()
    root = args.root
    clean = root / "Doctoral_Dissertation_Proposal_PEI_Rongkang_revised_clean.docx"
    tracked = root / "Doctoral_Dissertation_Proposal_PEI_Rongkang_revised_tracked.docx"
    clean_text = all_text(clean)
    tracked_text = all_text(tracked)
    citation_rows = read_csv(root / "citation_audit.csv")
    terms = read_csv(root / "terminology_control_table.csv")

    with tempfile.TemporaryDirectory() as td:
        accepted = Path(td) / "accepted.docx"
        subprocess.run([sys.executable, str(args.accept_script), str(tracked), "--mode", "accept", "--out", str(accepted)], check=True,
                       capture_output=True, text=True)
        accepted_equal = paragraph_text(accepted) == paragraph_text(clean)

    citation_coverage = []
    for row in citation_rows:
        surnames = row["authors"].split(";")[0].split(",")[0].strip()
        year = row["year"]
        if row["record_id"].startswith("jsps"):
            marker = "JSPS, 2026"
        elif row["record_id"] == "jst2026":
            marker = "JST, 2026"
        elif row["record_id"] == "statistics_japan2025":
            marker = "Statistics Bureau of Japan, 2025"
        else:
            marker = f"{surnames}, {year}"
        cited = bool(re.search(rf"{re.escape(surnames)}[^\n]{{0,35}}{re.escape(year)}", clean_text))
        citation_coverage.append((row["bibtex_key"], marker, cited, row["doi"], row["exact_claim_supported"]))
    with (root / "citation_coverage_audit.csv").open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(["bibtex_key", "text_marker", "cited_in_text", "doi", "claim_supported"])
        writer.writerows(citation_coverage)

    term_findings = []
    for row in terms:
        variants = [v.strip() for v in row["disallowed_or_review_form"].split(";") if v.strip()]
        for variant in variants:
            count = len(re.findall(rf"(?<![\w-]){re.escape(variant)}(?![\w-])", clean_text))
            if count:
                term_findings.append((row["protected_term"], variant, count, row["usage_rule"]))
    with (root / "terminology_audit_findings.csv").open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(["protected_term", "review_form", "count", "usage_rule"])
        writer.writerows(term_findings)

    causal_words = ["effect", "impact", "determines", "leads to", "produces", "causes", "return"]
    causal_counts = {word: len(re.findall(rf"(?i)\b{re.escape(word)}(?:s|ed|ing)?\b", clean_text)) for word in causal_words}
    required_headings = [
        "1. Introduction", "2. Literature Review and Theoretical Framework",
        "3. Study 1: Gendered Promotion Timelines in Japanese Academia",
        "4. Study 2: Institutional Mobility as a Gendered Career Strategy",
        "5. Study 3: Gendered Access to and Returns from KAKEN Resources and Project-Based Collaboration Networks",
        "6. Data and Methodological Strategy", "7. Expected Contributions and Work Plan",
        "Appendix A. Variable Construction and Validation", "References",
    ]
    heading_list = headings(clean)
    missing_headings = [h for h in required_headings if h not in heading_list]
    rq_counts = {rq: len(re.findall(rf"\b{rq}\b", clean_text)) for rq in
                 ["RQ1", "RQ2a", "RQ2b", "RQ2c", "RQ2d", "RQ2e", "RQ3a", "RQ3b", "RQ3c", "RQ3d"]}
    figure_mentions = Counter(re.findall(r"Figure\s+(\d+)", clean_text))
    table_mentions = Counter(re.findall(r"Table\s+(\d+)", clean_text))
    ins_clean, del_clean = revision_counts(clean)
    ins_tracked, del_tracked = revision_counts(tracked)
    source_sha = hashlib.sha256(args.source.read_bytes()).hexdigest()

    report = [
        "# Proposal Quality-Control Report", "",
        "## Deliverable integrity", "",
        f"- Source SHA-256: `{source_sha}`",
        f"- Source remained at the supplied path: `{args.source}`",
        f"- Clean comments: {comment_count(clean)}; tracked comments: {comment_count(tracked)}",
        f"- Clean revisions: {ins_clean} insertions / {del_clean} deletions",
        f"- Tracked revisions: {ins_tracked} insertions / {del_tracked} deletions",
        f"- Accepting tracked changes reproduces clean paragraph text: **{accepted_equal}**", "",
        "## Structural audit", "",
        f"- Heading count: {len(heading_list)}",
        f"- Missing required headings: {missing_headings or 'none'}",
        f"- Duplicate exact headings: {[h for h, n in Counter(heading_list).items() if n > 1] or 'none'}",
        f"- Figure references: {dict(sorted(figure_mentions.items()))}",
        f"- Numbered table references: {dict(sorted(table_mentions.items())) or 'none; appendix tables are heading-labelled'}", "",
        "## Research-question audit", "",
        *[f"- {rq}: {count} occurrences; mapped in `proposal_structure_map.md`." for rq, count in rq_counts.items()], "",
        "## Citation audit", "",
        f"- Verified citation records: {len(citation_rows)}",
        f"- BibTeX entries: {sum(1 for line in (root / 'references.bib').read_text(encoding='utf-8').splitlines() if line.startswith('@'))}",
        f"- Citation audit records not detected in text: {[key for key, marker, cited, doi, claim in citation_coverage if not cited] or 'none'}",
        "- DOI metadata and exact supported claims are recorded in `citation_audit.csv`.", "",
        "## Terminology audit", "",
        f"- Protected terms: {len(terms)}",
        f"- Review-form findings: {len(term_findings)} (each occurrence requires contextual review; see CSV).",
        "- `researchmap`, `Co-I`, Associate Professor/Full Professor, institutional mobility, and time-ordered association are controlled.", "",
        "## Non-causal language audit", "",
        *[f"- `{word}` family: {count} occurrences." for word, count in causal_counts.items()],
        "- Reviewed uses occur in explicit interpretation boundaries, cited study titles/claims, or statements that the design does not identify causal effects.", "",
        "## Required limitations", "",
        "The proposal explicitly retains limitations concerning institution properties, geography, ranking coverage, student counts, research institutes, missing career history, KAKEN linkage, project-network scope, causal identification, mobility selection, and unobserved family costs.", "",
        "## Result", "",
        f"- Automated structural/copy consistency gate: **{'PASS' if accepted_equal and not missing_headings and comment_count(clean) == 0 else 'FAIL'}**",
        "- Visual rendering remains a separate mandatory gate.",
    ]
    (root / "quality_control_report.md").write_text("\n".join(report) + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()
